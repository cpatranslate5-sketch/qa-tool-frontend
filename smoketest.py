"""Quick local smoke test for the shared-folders / admin model, the
structured tone-of-address document, and the new check types. Uses
a throwaway SQLite DB (no DATABASE_URL set) and no AI key, so only
rule-based findings are expected, not AI ones."""
import asyncio
import io
import os
import sys

os.environ["DATABASE_URL"] = ""

sys.path.insert(0, os.path.dirname(__file__))

import app.database as dbmod
dbmod.engine = dbmod.create_engine("sqlite:///./qa_tool_smoketest.db", connect_args={"check_same_thread": False})
dbmod.SessionLocal = dbmod.sessionmaker(autocommit=False, autoflush=False, bind=dbmod.engine)
if os.path.exists("qa_tool_smoketest.db"):
    os.remove("qa_tool_smoketest.db")
dbmod.init_db()

from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

def check(label, resp, expect=200):
    status = "OK" if resp.status_code == expect else "FAIL"
    print(f"[{status}] {label} -> {resp.status_code}")
    if status == "FAIL":
        print("   body:", resp.text[:500])
    return resp

# --- folder creation: first ever becomes admin ---
r = check("create folder Александр (first -> admin)", client.post("/managers", json={"name": "Александр", "code": "1234"}))
admin_id = r.json()["id"]
assert r.json()["is_admin"] is True

r = check("create folder Мария (second -> not admin)", client.post("/managers", json={"name": "Мария", "code": "5678"}))
regular_id = r.json()["id"]
assert r.json()["is_admin"] is False

check("duplicate folder name", client.post("/managers", json={"name": "Александр", "code": "0000"}), expect=409)

# --- listing folders (public, no code) — admin folder always first
# (point 2 of Александр's folder-access spec) ---
r = check("list managers", client.get("/managers"))
assert len(r.json()) == 2
assert r.json()[0]["is_admin"] is True and r.json()[0]["name"] == "Александр"

# --- unlock flow ---
check("unlock wrong code", client.post(f"/managers/{regular_id}/unlock", json={"code": "wrong"}), expect=401)
check("unlock correct code", client.post(f"/managers/{regular_id}/unlock", json={"code": "5678"}))

# --- admin bypass: whoever already has admin access on this device can
# open any other folder without typing that folder's own password ---
check("admin-enter refuses a non-admin claimed id", client.post(
    f"/managers/{regular_id}/admin-enter", json={"admin_manager_id": regular_id}
), expect=403)
r = check("real admin opens Мария's folder with no password", client.post(
    f"/managers/{regular_id}/admin-enter", json={"admin_manager_id": admin_id}
))
assert r.json()["id"] == regular_id and r.json()["name"] == "Мария"

# --- password change: available to every folder, not just admin ---
check("wrong current password rejected", client.post(
    f"/managers/{regular_id}/change-password", json={"current_code": "nope", "new_code": "9999"}
), expect=401)
check("password change succeeds", client.post(
    f"/managers/{regular_id}/change-password", json={"current_code": "5678", "new_code": "9999"}
))
check("old password no longer works", client.post(f"/managers/{regular_id}/unlock", json={"code": "5678"}), expect=401)
check("new password works", client.post(f"/managers/{regular_id}/unlock", json={"code": "9999"}))

# --- non-admin blocked from structural changes ---
check("non-admin create project blocked", client.post("/projects", json={"name": "Pragmatic Play Promo", "manager_id": regular_id}), expect=403)

# --- admin creates project ---
r = check("admin create project", client.post("/projects", json={"name": "Pragmatic Play Promo", "manager_id": admin_id}))
project_id = r.json()["id"]
assert r.json()["created_by_name"] == "Александр"
assert r.json()["tone_filename"] == ""

check("duplicate project name", client.post("/projects", json={"name": "Pragmatic Play Promo", "manager_id": admin_id}), expect=409)

# --- no more language folders: everything runs directly against the project ---
import openpyxl

# --- doc-gating: tone (register) check refuses to run until its doc exists ---
check("tone (register) check blocked with no doc uploaded", client.post("/check", json={
    "source": "Play now.",
    "translation": "Играйте сейчас.",
    "checks": ["register"],
    "project_id": project_id,
    "source_lang": "en",
    "target_lang": "ru",
    "manager_name": "Мария",
    "manager_id": regular_id,
}), expect=400)

# --- Tone-of-address doc: real layout, matching the agency's actual
# export — language codes as header-row columns, not one row per language
# as originally assumed, with the register in the data row below each
# column. "KZ" is the same country-code-style label the real file uses
# for Kazakh.
twb = openpyxl.Workbook()
tws = twb.active
tws.append(["EN", "RU", "KZ", "ES (MX)"])
tws.append(["Формальное", "Формальное", "Формальное", "Неформальное"])
tbuf = io.BytesIO()
twb.save(tbuf)
tbuf.seek(0)
r = check("admin tone upload", client.post(
    f"/projects/{project_id}/tone/upload",
    files={"file": ("tone.xlsx", tbuf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"manager_id": admin_id},
))
assert r.json()["rule_count"] == 4, r.json()

# --- known languages (the checkbox catalog) is now its OWN list,
# completely decoupled from the tone-of-address document — Александр
# asked for the checkbox catalog to change ONLY on an explicit
# add/remove, never as a side effect of uploading anything. Right after
# a tone upload, the catalog is still empty: uploading a tone doc alone
# no longer populates it. ---
r = check("known languages is NOT auto-populated by a tone upload", client.get(f"/projects/{project_id}/known-languages"))
assert r.json()["languages"] == [], r.json()

# --- admin builds the catalog explicitly, one language at a time ---
check("non-admin can't add a catalog language", client.post(
    f"/projects/{project_id}/languages", json={"manager_id": regular_id, "lang_code": "ru"}
), expect=403)
for code in ["ru", "es-mx", "kz", "en"]:
    check(f"admin adds '{code}' to the catalog", client.post(
        f"/projects/{project_id}/languages", json={"manager_id": admin_id, "lang_code": code}
    ))
# re-adding an already-present language is a harmless no-op, not an error
check("re-adding an existing catalog language is a no-op", client.post(
    f"/projects/{project_id}/languages", json={"manager_id": admin_id, "lang_code": "ru"}
))
r = check("known languages now reflects the manually-built catalog", client.get(f"/projects/{project_id}/known-languages"))
assert set(r.json()["languages"]) == {"ru", "es-mx", "kz", "en"}, r.json()

# --- admin can drop a single straggler language from the catalog without
# touching the rest — for exactly the situation this feature was built
# for: a project created via "copy from an existing project" (tested
# further below) inherits that other project's whole catalog, including
# a language nobody meant for THIS project ---
check("non-admin can't delete a catalog language", client.delete(
    f"/projects/{project_id}/languages/kz", params={"manager_id": regular_id}
), expect=403)
check("deleting a language not in the catalog 404s", client.delete(
    f"/projects/{project_id}/languages/zz", params={"manager_id": admin_id}
), expect=404)
# uppercase on the way in, to prove the match is case-insensitive (the
# frontend always displays codes upper-cased) even though it's stored
# lower-cased
r = check("admin deletes the stray 'kz' catalog language", client.delete(
    f"/projects/{project_id}/languages/KZ", params={"manager_id": admin_id}
))
assert set(r.json()["languages"]) == {"ru", "es-mx", "en"}, r.json()
r = check("known languages no longer include the deleted one", client.get(f"/projects/{project_id}/known-languages"))
assert set(r.json()["languages"]) == {"ru", "es-mx", "en"}, r.json()

# --- re-uploading the tone doc doesn't touch the catalog either (fully
# decoupled in both directions) — rule_count changes, known-languages
# doesn't ---
tone_recheck_wb = openpyxl.Workbook()
tone_recheck_ws = tone_recheck_wb.active
tone_recheck_ws.append(["EN", "RU", "JA"])
tone_recheck_ws.append(["Формальное", "Формальное", "Формальное"])
tone_recheck_buf = io.BytesIO()
tone_recheck_wb.save(tone_recheck_buf)
tone_recheck_buf.seek(0)
r = check("re-uploading the tone doc leaves the catalog untouched", client.post(
    f"/projects/{project_id}/tone/upload",
    files={"file": ("tone_reupload_check.xlsx", tone_recheck_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"manager_id": admin_id},
))
assert r.json()["rule_count"] == 3, r.json()  # the tone doc itself DID change...
r = check("catalog unchanged after tone re-upload", client.get(f"/projects/{project_id}/known-languages"))
assert set(r.json()["languages"]) == {"ru", "es-mx", "en"}, r.json()  # ...but the catalog didn't

# --- now that the tone doc exists, the previously-blocked check runs fine ---
check("tone check now runs", client.post("/check", json={
    "source": "Play now.",
    "translation": "Играйте сейчас.",
    "checks": ["register"],
    "project_id": project_id,
    "source_lang": "en",
    "target_lang": "ru",
    "manager_name": "Мария",
    "manager_id": regular_id,
}))

# --- BOTH folders see the same shared project (no manager scoping) ---
r = check("list projects (shared)", client.get("/projects"))
assert len(r.json()) == 1

# --- non-admin CAN run a single check across every remaining check type ---
r = check("non-admin single check (ru)", client.post("/check", json={
    "source": "The bonus is $50 and expires in 3 days.",
    "translation": "Бонус составляет $500 и истекает через 3 дня",
    "checks": ["numbers", "placeholders", "register", "typo", "untranslatable", "completeness", "punctuation"],
    "project_id": project_id,
    "source_lang": "en",
    "target_lang": "ru",
    "manager_name": "Мария",
    "manager_id": regular_id,
}))
findings = r.json()["findings"]
print("   findings:", findings)
# number mismatch ($50 vs $500) must be caught by the rule-based check
assert any(f["type"] == "numbers" for f in findings)
# source ends in "." and translation doesn't -> punctuation rule should fire
assert any(f["type"] == "punctuation" for f in findings), findings

# --- double space is caught regardless of source ---
r = check("punctuation: double space", client.post("/check", json={
    "source": "Hello world.",
    "translation": "Привет  мир.",
    "checks": ["punctuation"],
    "manager_name": "Мария",
}))
findings = r.json()["findings"]
assert any(f["type"] == "punctuation" for f in findings), findings

# --- extra_instructions field is accepted and doesn't break anything ---
check("check with extra_instructions", client.post("/check", json={
    "source": "Play Golden Spin now!",
    "translation": "Играйте в Golden Spin сейчас!",
    "checks": ["untranslatable"],
    "project_id": project_id,
    "source_lang": "en",
    "target_lang": "ru",
    "extra_instructions": "В этой задаче 'Golden Spin' нужно переводить как 'Голден Спин'.",
    "manager_name": "Мария",
    "manager_id": regular_id,
}))

# --- history shows who performed it, keyed by project + folder (each
# manager only sees their own runs — point 1 of Александр's spec) ---
r = check("history shows attribution", client.get(f"/projects/{project_id}/history", params={"manager_id": regular_id}))
history = r.json()
assert len(history) == 3  # tone + full-checks + extra_instructions check (double-space was standalone)
assert history[0]["performed_by_name"] == "Мария"
assert history[0]["source_lang"] == "en" and history[0]["target_lang"] == "ru"
print("   performed_by_name:", history[0]["performed_by_name"])

# --- history is scoped per folder, not shared across every folder that
# touches the project ---
check("admin runs a check on the same shared project", client.post("/check", json={
    "source": "Hello.",
    "translation": "Привет.",
    "checks": ["punctuation"],
    "project_id": project_id,
    "source_lang": "en",
    "target_lang": "ru",
    "manager_name": "Александр",
    "manager_id": admin_id,
}))
r = check("Мария's history unaffected by admin's check", client.get(f"/projects/{project_id}/history", params={"manager_id": regular_id}))
assert len(r.json()) == 3, r.json()
r = check("admin's own history shows only admin's check", client.get(f"/projects/{project_id}/history", params={"manager_id": admin_id}))
assert len(r.json()) == 1, r.json()
assert r.json()[0]["performed_by_name"] == "Александр"

# --- deleting a single (point) check from history ---
single_check_to_delete = history[0]["id"]
check("admin can't delete Мария's single check (not theirs)", client.delete(
    f"/projects/{project_id}/history/{single_check_to_delete}", params={"manager_id": admin_id}
), expect=404)
check("Мария can delete her own single check", client.delete(
    f"/projects/{project_id}/history/{single_check_to_delete}", params={"manager_id": regular_id}
))
r = check("deleted single check no longer in history", client.get(
    f"/projects/{project_id}/history", params={"manager_id": regular_id}
))
assert len(r.json()) == 2, r.json()
assert all(h["id"] != single_check_to_delete for h in r.json()), r.json()

# --- non-admin CAN run a multi-check upload ---
sample_path = "/root/.claude/uploads/aee9e6e5-e96f-5b4b-aa4e-8aad6284c8c9/147efc1b-Promo_Rules_Localization.xlsx"
with open(sample_path, "rb") as f:
    r = check("non-admin multi-check upload", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": ""},
    ))
multi_data = r.json()
multi_check_id = multi_data["multi_check_id"]
print("   summary:", multi_data["summary"])
# This sample file is above BATCH_THRESHOLD_CHARS (many languages), so it
# would normally go through the Message Batches path — but with no API key
# configured there's nothing to submit, so it finalizes immediately with
# just the rule-based findings, same as the old fully-synchronous behavior.
assert multi_data["status"] == "completed", multi_data
# --- a completed check reports when it started/finished, so the report
# (and history list) can show how long it actually took ---
assert multi_data["created_at"], multi_data
assert multi_data["completed_at"], multi_data
print("[OK] completed multi-check response includes created_at/completed_at")

# --- extend the catalog to a more realistic size before exercising
# detect-languages against the real sample file (which spans ~30
# languages) — mirrors an admin gradually building out their real
# language list over time, on top of the small set used above to test
# plain add/remove ---
for code in ["ar", "kk", "pt-br"]:
    check(f"admin adds '{code}' to the catalog", client.post(
        f"/projects/{project_id}/languages", json={"manager_id": admin_id, "lang_code": code}
    ))

# --- detect-languages: reports which of the file's language-shaped
# columns match the project's OWN catalog (checkable) vs. which merely
# LOOK like a language code but aren't on the manager's list at all
# (unknown_languages) — the fix for Александр's concrete bug report: a
# column literally labelled "PR" (meant as an abbreviation for
# Portuguese, but not a real code for it) used to get silently treated
# as a real target language, with Peru's flag, purely because it parsed
# as language-shaped. Now nothing enters the checkbox list just because
# a file happens to contain it. ---
with open(sample_path, "rb") as f:
    r = check("detect-languages splits catalog-matched vs unknown languages", client.post(
        f"/projects/{project_id}/multi-check/detect-languages",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ))
detected = r.json()["languages"]
unknown_detected = r.json()["unknown_languages"]
# every catalog language actually present in the file is reported as a
# real, checkable target...
assert {"ru", "es-mx", "en", "ar", "kk", "pt-br"} <= set(detected), detected
# ...while a real language column the manager simply hasn't added to
# their catalog yet (Bengali) is reported separately, not silently mixed
# into the checkable list
assert "bn" in unknown_detected, unknown_detected
assert r.json()["unrecognized_columns"] == [], r.json()  # this sample file has none of those
print(f"   detected (catalog-matched) languages: {detected}")
print(f"   unknown (language-shaped but not on the catalog) languages: {unknown_detected}")

# --- the exact scenario Александр reported: a column literally labelled
# "PR" (a manager's mistaken abbreviation for Portuguese) must come back
# as unknown — never silently added as a real target language ---
mislabel_wb = openpyxl.Workbook()
mislabel_ws = mislabel_wb.active
mislabel_ws.append(["Context", "en", "ru", "PR"])
mislabel_ws.append(["Greeting", "Hello", "Привет", "Olá"])
mislabel_buf = io.BytesIO()
mislabel_wb.save(mislabel_buf)
mislabel_buf.seek(0)
r = check("a mislabeled column ('PR' for Portuguese) is reported as unknown, not added silently", client.post(
    f"/projects/{project_id}/multi-check/detect-languages",
    files={"file": ("mislabeled_pr.xlsx", mislabel_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
))
assert "pr" not in r.json()["languages"], r.json()
assert "pr" in r.json()["unknown_languages"], r.json()
assert "en" in r.json()["languages"] and "ru" in r.json()["languages"], r.json()

# --- exactly the fix Александр asked for: once the manager explicitly
# adds the (genuinely new) language to the catalog, the SAME file is
# re-checked and that column now counts as a recognized target — never
# automatically, only after the explicit add ---
check("admin adds the genuinely-new 'pr' language to the catalog", client.post(
    f"/projects/{project_id}/languages", json={"manager_id": admin_id, "lang_code": "pr"}
))
mislabel_buf.seek(0)
r = check("after adding it to the catalog, the same column is now recognized", client.post(
    f"/projects/{project_id}/multi-check/detect-languages",
    files={"file": ("mislabeled_pr.xlsx", mislabel_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
))
assert "pr" in r.json()["languages"], r.json()
assert "pr" not in r.json()["unknown_languages"], r.json()
# clean up so 'pr' doesn't leak into later known-languages assertions
check("clean up: remove 'pr' from the catalog again", client.delete(
    f"/projects/{project_id}/languages/pr", params={"manager_id": admin_id}
))

# --- the Portuguese default: a column literally labelled bare "PT" (no
# region at all) must resolve straight to "pt-br", matching the "pt-br"
# already on this project's catalog (added above) — not show up as
# "unknown", and not stay bare "pt" either ---
pt_wb = openpyxl.Workbook()
pt_ws = pt_wb.active
pt_ws.append(["Context", "en", "ru", "PT"])
pt_ws.append(["Greeting", "Hello", "Привет", "Olá"])
pt_buf = io.BytesIO()
pt_wb.save(pt_buf)
pt_buf.seek(0)
r = check("a bare 'PT' column resolves to Brazilian Portuguese ('pt-br'), matching the catalog", client.post(
    f"/projects/{project_id}/multi-check/detect-languages",
    files={"file": ("bare_pt.xlsx", pt_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
))
assert "pt-br" in r.json()["languages"], r.json()
assert "pt" not in r.json()["languages"], r.json()
assert "pt" not in r.json()["unknown_languages"] and "pt-br" not in r.json()["unknown_languages"], r.json()

# --- a column that isn't recognized as a language must be reported back
# BEFORE the manager presses "start", not only inside a finished report —
# by which point an AI-backed check may already have run without ever
# covering a genuine language column that got missed. ---
unrec_wb = openpyxl.Workbook()
unrec_ws = unrec_wb.active
unrec_ws.append(["Context", "en", "ru", "Notes for reviewer"])
unrec_ws.append(["Greeting", "Hello", "Привет", "double-check tone"])
unrec_buf = io.BytesIO()
unrec_wb.save(unrec_buf)
unrec_buf.seek(0)
r = check("detect-languages also reports unrecognized columns up front", client.post(
    f"/projects/{project_id}/multi-check/detect-languages",
    files={"file": ("with_notes_column.xlsx", unrec_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
))
assert "Notes for reviewer" in r.json()["unrecognized_columns"], r.json()
assert "en" in r.json()["languages"] and "ru" in r.json()["languages"], r.json()

# --- Александр hit this live right after the above shipped: his real file
# spells some region-qualified languages as "ES MX"/"PT BR" — the display
# style ("ES (MX)") without the parentheses — which fell through to the
# generic "looks like prose" rejection (any header with a space in it) and
# was reported as an unrecognized column even though it very much is a
# language. Deliberately ALL CAPS only, so an ordinary two-word column like
# "Task name" (mixed case in every real file) still correctly stays
# unrecognized rather than being guessed as a language. ---
space_wb = openpyxl.Workbook()
space_ws = space_wb.active
space_ws.append(["Context", "en", "ES MX", "PT BR", "Task name"])
space_ws.append(["Greeting", "Hello", "Hola", "Ola", "internal note"])
space_buf = io.BytesIO()
space_wb.save(space_buf)
space_buf.seek(0)
r = check("detect-languages recognizes ALL-CAPS \"ES MX\"/\"PT BR\"-style headers as languages", client.post(
    f"/projects/{project_id}/multi-check/detect-languages",
    files={"file": ("space_style_codes.xlsx", space_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
))
assert "es-mx" in r.json()["languages"] and "pt-br" in r.json()["languages"], r.json()
assert "Task name" in r.json()["unrecognized_columns"], r.json()
assert "ES MX" not in r.json()["unrecognized_columns"], r.json()

# --- Александр's redesign: instead of the manager reviewing an
# auto-generated "here's what we found" list (easy to miss an absence
# from), the manager states which languages they expect up front, and
# verify-languages is the explicit per-language yes/no this drives —
# found via the same safe bridging as everywhere else (a code spelled
# differently still counts), missing only when nothing safely matches. ---
with open(sample_path, "rb") as f:
    r = check("verify-languages reports found/not-found per requested language", client.post(
        f"/projects/{project_id}/multi-check/verify-languages",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"codes": "ru, es-mx, nl"},  # nl (Dutch) genuinely isn't in this sample file
    ))
verify_results = {row["code"]: row["found"] for row in r.json()["results"]}
assert verify_results == {"ru": True, "es-mx": True, "nl": False}, verify_results
# and the space-style header fix above is itself confirmed found through
# this same endpoint, not just through detect-languages.
verify_space_buf = io.BytesIO()
space_wb.save(verify_space_buf)
verify_space_buf.seek(0)
r = check("verify-languages also recognizes the ALL-CAPS space-style headers", client.post(
    f"/projects/{project_id}/multi-check/verify-languages",
    files={"file": ("space_style_codes.xlsx", verify_space_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"codes": "es-mx, pt-br, de"},  # de genuinely isn't in this file
))
verify_results2 = {row["code"]: row["found"] for row in r.json()["results"]}
assert verify_results2 == {"es-mx": True, "pt-br": True, "de": False}, verify_results2
print("[OK] verify-languages: explicit per-language found/not-found confirmation, driving "
      "the new \"tick what you expect, confirm, get told exactly what's missing\" flow")

# --- target_langs filter: checking just 2 of the file's many languages
# should only touch those 2 in the summary ---
with open(sample_path, "rb") as f:
    r = check("multi-check with target_langs filter", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": "", "target_langs": "ru,es-mx"},
    ))
filtered_data = r.json()
assert filtered_data["status"] == "completed", filtered_data
assert set(filtered_data["summary"]["languages_checked"]) == {"ru", "es-mx"}, filtered_data["summary"]
print("   filtered languages_checked:", filtered_data["summary"]["languages_checked"])

r = check("multi-check history shows attribution", client.get(
    f"/projects/{project_id}/multi-check", params={"manager_id": regular_id}
))
assert r.json()[0]["performed_by_name"] == "Мария"
assert r.json()[0]["status"] == "completed"

report_resp = check("multi-check report download", client.get(
    f"/projects/{project_id}/multi-check/{multi_check_id}/report.xlsx", params={"manager_id": regular_id}
))
# --- the report's first row should say how long the check took, now that
# Александр asked for the time spent to show up in the downloadable report
# too (not just the history list) ---
report_wb = openpyxl.load_workbook(io.BytesIO(report_resp.content))
report_ws = report_wb.active
report_first_cell = report_ws.cell(row=1, column=1).value
assert "заняла" in report_first_cell, report_first_cell
assert report_ws.cell(row=3, column=1).value == "Лист", "header row should follow the duration line + spacer"
print("[OK] downloaded report's first row states check duration:", report_first_cell)

# --- multi-check history/detail/report are also scoped per folder ---
r = check("admin's multi-check history is empty (Мария's uploads aren't his)", client.get(
    f"/projects/{project_id}/multi-check", params={"manager_id": admin_id}
))
assert r.json() == [], r.json()
check("admin can't fetch Мария's multi-check detail by id", client.get(
    f"/projects/{project_id}/multi-check/{multi_check_id}", params={"manager_id": admin_id}
), expect=404)
check("admin can't download Мария's multi-check report", client.get(
    f"/projects/{project_id}/multi-check/{multi_check_id}/report.xlsx", params={"manager_id": admin_id}
), expect=404)

# --- regression: exactly what Александр hit live. His catalog checkbox
# for Portuguese was ticked as a bare "pt" (added before Portuguese
# started defaulting to Brazilian), but the file's own "PT" column now
# normalizes to "pt-br" — a plain string-equality filter would silently
# drop it from the actual check even though detect-languages had already
# told him it would be checked. The real run must bridge this the same
# way detect-languages does, never drop a language silently. ---
pt_check_wb = openpyxl.Workbook()
pt_check_ws = pt_check_wb.active
pt_check_ws.append(["Context", "en", "PT"])
pt_check_ws.append(["Greeting", "Hello", "Olá"])
pt_check_buf = io.BytesIO()
pt_check_wb.save(pt_check_buf)
pt_check_buf.seek(0)
r = check(
    "multi-check with an older, coarser target_langs spelling ('pt') still checks the file's actual 'pt-br' column",
    client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("bare_pt_check.xlsx", pt_check_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "en", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": "", "target_langs": "pt"},
    ),
)
pt_check_data = r.json()
assert pt_check_data["status"] == "completed", pt_check_data
assert "pt-br" in pt_check_data["summary"]["languages_checked"], pt_check_data["summary"]
assert "pt-br" in pt_check_data["sheets"][0]["languages"], pt_check_data["sheets"][0]

# --- a "DO NOT TRANSLATE" cell means this row is deliberately left
# untranslated for this language on purpose — it must be skipped entirely
# (no finding at all), not flagged as a numbers/content mismatch, even
# though the literal text would otherwise clearly disagree with the
# source. Александр's files legitimately need some rows translated for one
# language but not another. ---
dnt_wb = openpyxl.Workbook()
dnt_ws = dnt_wb.active
dnt_ws.append(["EN", "RU"])
dnt_ws.append(["Price: 500 dollars.", "Цена: 400 долларов."])  # genuine mismatch — must still be caught
dnt_ws.append(["Price: 500 dollars.", "DO NOT TRANSLATE"])
dnt_ws.append(["Price: 500 dollars.", " [do NOT Translate] "])  # brackets + mixed case variant
dnt_buf = io.BytesIO()
dnt_wb.save(dnt_buf)
dnt_buf.seek(0)
r = check("multi-check with DO NOT TRANSLATE cells", client.post(
    f"/projects/{project_id}/multi-check",
    files={"file": ("dnt.xlsx", dnt_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"source_lang": "en", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": "", "checks": "numbers"},
))
dnt_data = r.json()
assert dnt_data["status"] == "completed", dnt_data
dnt_ru_rows = dnt_data["sheets"][0]["languages"]["ru"]
assert len(dnt_ru_rows) == 1, dnt_ru_rows  # only the genuine mismatch row, both DNT rows skipped entirely
assert dnt_ru_rows[0]["translation"] == "Цена: 400 долларов.", dnt_ru_rows
assert any(f["type"] == "numbers" for f in dnt_ru_rows[0]["findings"]), dnt_ru_rows
print("   DO NOT TRANSLATE rows correctly skipped, genuine mismatch still caught:", dnt_ru_rows)

# --- Александр asked "стоимость 0$ с найденными проблемами — это
# нормально?" — yes, when only the free algorithmic checks are selected
# (no AI call ever happens, so nothing to bill), and checks_run in the
# response is exactly what lets the UI show which criteria actually ran
# instead of leaving him guessing. Covers both at once: run with only the
# free "punctuation" check (which is one of the CHECK_OPTIONS-labeled keys,
# unlike the plain "numbers"/"max_length" the frontend silently folds in),
# confirm cost stays 0 and checks_run reflects exactly that selection. ---
only_algo_wb = openpyxl.Workbook()
only_algo_ws = only_algo_wb.active
only_algo_ws.append(["EN", "RU"])
only_algo_ws.append(["Hello world.", "Привет мир"])  # source ends with "." but translation has no terminal punctuation at all -> reliably flagged
only_algo_buf = io.BytesIO()
only_algo_wb.save(only_algo_buf)
only_algo_buf.seek(0)
r = check("multi-check with only a free algorithmic criterion selected", client.post(
    f"/projects/{project_id}/multi-check",
    files={"file": ("only-algo.xlsx", only_algo_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"source_lang": "en", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": "", "checks": "punctuation"},
))
only_algo_data = r.json()
assert only_algo_data["status"] == "completed", only_algo_data
assert only_algo_data["cost_usd"] == 0, only_algo_data  # no AI check type was even selected
assert only_algo_data["checks_run"] == ["punctuation"], only_algo_data
print("[OK] $0 cost with real findings is expected when only free/algorithmic criteria are selected; checks_run confirms which criteria ran")

# --- the downloadable report should let Александр filter by language using
# Excel's own column-filter control ("фильтр по языкам") ---
only_algo_report = check("report for the only-algo check downloads fine", client.get(
    f"/projects/{project_id}/multi-check/{only_algo_data['multi_check_id']}/report.xlsx", params={"manager_id": regular_id}
))
only_algo_wb_read = openpyxl.load_workbook(io.BytesIO(only_algo_report.content))
only_algo_ws_read = only_algo_wb_read.active
# The filter range must start exactly ON the header row ("Лист", "Строка в
# файле", ...), not one row early on the blank spacer above it — an
# earlier version of this got that boundary wrong and silently excluded
# the header (and included the empty spacer instead) from the filterable
# range.
header_row_idx = next(
    r for r in range(1, only_algo_ws_read.max_row + 1)
    if only_algo_ws_read.cell(row=r, column=1).value == "Лист"
)
assert only_algo_ws_read.auto_filter.ref == f"A{header_row_idx}:I{only_algo_ws_read.max_row}", (
    only_algo_ws_read.auto_filter.ref, header_row_idx, only_algo_ws_read.max_row
)
print("[OK] downloaded report's Excel column filter starts exactly on the header row:", only_algo_ws_read.auto_filter.ref)

# --- deleting a multi-check report/upload from history ---
check("Мария can delete her own multi-check", client.delete(
    f"/projects/{project_id}/multi-check/{multi_check_id}", params={"manager_id": regular_id}
))
r = check("deleted multi-check no longer in history", client.get(
    f"/projects/{project_id}/multi-check", params={"manager_id": regular_id}
))
assert all(h["id"] != multi_check_id for h in r.json()), r.json()
check("deleted multi-check detail is gone", client.get(
    f"/projects/{project_id}/multi-check/{multi_check_id}", params={"manager_id": regular_id}
), expect=404)
check("admin can't delete Мария's multi-check (not theirs)", client.delete(
    f"/projects/{project_id}/multi-check/{dnt_data['multi_check_id']}", params={"manager_id": admin_id}
), expect=404)

# --- large multi-check actually goes through the Message Batches path when
# a batch can be submitted — simulate that here (no real Anthropic key in
# this sandbox) by faking the three network calls, to prove the
# submit -> poll -> merge-AI-findings-into-the-rule-based-skeleton pipeline
# actually produces a correct, complete result. ---
import app.excel_multi as excel_multi_mod

# custom_id is deliberately opaque (see build_batch_plan — it's built from
# the sheet/position index only, never from the language code itself, so a
# weird character in a file's own language column can never produce an
# invalid custom_id and get the whole batch rejected by Anthropic). This
# fake captures whatever custom_ids were actually submitted rather than
# hardcoding the old "s{sheet}-{lang}" shape, so the test doesn't silently
# stop verifying anything if that internal scheme ever changes again.
_submitted_custom_ids: list[str] = []


async def _fake_create_message_batch(requests):
    assert requests, "expected at least one per-language batch request to be built"
    _submitted_custom_ids.clear()
    _submitted_custom_ids.extend(r["custom_id"] for r in requests)
    return "msgbatch_test123"


_batch_status_call_count = {"n": 0}


async def _fake_get_batch_status(batch_id):
    assert batch_id == "msgbatch_test123"
    _batch_status_call_count["n"] += 1
    if _batch_status_call_count["n"] == 1:
        # First check (triggered by the history list's own opportunistic
        # poll, below) — still running, so it stays "processing" with real
        # counts instead of finalizing right there.
        return {"processing_status": "in_progress", "request_counts": {
            "processing": 5, "succeeded": 3, "errored": 0, "canceled": 0, "expired": 0,
        }}
    return {"processing_status": "ended", "results_url": "fake://results"}


async def _fake_get_batch_results(results_url):
    assert results_url == "fake://results"
    # Same fake AI finding for every custom_id that was actually submitted —
    # shape matches the real get_batch_results: {custom_id: {"text": ...,
    # "usage": ..., "stop_reason": ..., "result_type": ...}}. Every
    # language's rows are identical in this sample file's languages, so
    # this reliably shows up under "ru" (and every other language) without
    # the test needing to know its exact custom_id.
    assert _submitted_custom_ids, "expected _fake_create_message_batch to have run first"
    return {
        cid: {
            "text": '[{"row": 1, "type": "typo", "severity": "medium", "message": "тестовая ИИ-находка"}]',
            "usage": {"input_tokens": 1000, "output_tokens": 200},
            "stop_reason": "end_turn",
            "result_type": "succeeded",
        }
        for cid in _submitted_custom_ids
    }


excel_multi_mod.create_message_batch = _fake_create_message_batch
excel_multi_mod.get_batch_status = _fake_get_batch_status
excel_multi_mod.get_batch_results = _fake_get_batch_results

with open(sample_path, "rb") as f:
    r = check("large multi-check submits as a batch", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": ""},
    ))
batch_multi_data = r.json()
assert batch_multi_data["status"] == "processing", batch_multi_data
batch_multi_check_id = batch_multi_data["multi_check_id"]
# Submission response already knows the total request count for free (no
# Anthropic call needed to say "0 of N so far").
assert batch_multi_data["progress"]["done"] == 0, batch_multi_data
assert batch_multi_data["progress"]["total"] > 0, batch_multi_data
# The UI falls back to showing elapsed waiting time whenever Anthropic's own
# counts haven't moved yet (Александр found the percentage looked frozen at
# 0% for a long stretch) — needs a real timestamp to compute that from.
assert batch_multi_data["created_at"], batch_multi_data

r = check("history shows the batch entry as processing", client.get(
    f"/projects/{project_id}/multi-check", params={"manager_id": regular_id}
))
hist_entry = r.json()[0]
assert hist_entry["status"] == "processing", hist_entry
# The history list opportunistically polls Anthropic itself (one call per
# still-processing upload) so a manager sees real progress — "готово X из
# Y" — without having to open that specific check first.
assert hist_entry["progress"] == {"done": 3, "total": 8}, hist_entry

check("report download blocked while processing", client.get(
    f"/projects/{project_id}/multi-check/{batch_multi_check_id}/report.xlsx", params={"manager_id": regular_id}
), expect=409)

# First poll: our fake get_batch_status already says "ended", so this same
# call both notices completion and merges the results in.
r = check("polling finalizes the batch", client.get(
    f"/projects/{project_id}/multi-check/{batch_multi_check_id}", params={"manager_id": regular_id}
))
finalized = r.json()
assert finalized["status"] == "completed", finalized
ru_findings = finalized["sheets"][0]["languages"].get("ru", [])
assert any(
    any(f["type"] == "typo" and "тестовая ИИ-находка" in f["message"] for f in row["findings"])
    for row in ru_findings
), ru_findings
print("   ru findings after batch merge:", ru_findings)
# The batch's real Anthropic usage (faked above) must translate into a
# non-zero cost, computed with the batch discount, and persisted on the
# record (not just present in the one-off response).
assert finalized["cost_usd"] > 0, finalized
# A batch-processed check finalizes asynchronously (unlike the synchronous
# path checked earlier), so completed_at has to be set separately, right
# here at finalization time — confirm that actually happened, not just for
# the synchronous path.
assert finalized["created_at"], finalized
assert finalized["completed_at"], finalized
print("[OK] finalized batch check also gets created_at/completed_at (not just the synchronous path)")
r2 = check("multi-check detail re-fetch still shows the persisted cost", client.get(
    f"/projects/{project_id}/multi-check/{batch_multi_check_id}", params={"manager_id": regular_id}
))
assert r2.json()["cost_usd"] == finalized["cost_usd"], r2.json()
print(f"   batch cost_usd: {finalized['cost_usd']}")

r = check("history now shows completed", client.get(
    f"/projects/{project_id}/multi-check", params={"manager_id": regular_id}
))
assert r.json()[0]["status"] == "completed"
assert r.json()[0]["completed_at"], r.json()[0]

check("report download works once completed", client.get(
    f"/projects/{project_id}/multi-check/{batch_multi_check_id}/report.xlsx", params={"manager_id": regular_id}
))

# --- while a batch is still "in_progress" (not yet "ended"), the detail
# endpoint surfaces Anthropic's own request_counts as {"done", "total"}
# instead of finalizing — this is what lets the UI show a real progress
# readout ("2 of 5 done") rather than a guessed time estimate, which
# Anthropic's API doesn't provide at all. ---
_progress_poll_count = {"n": 0}


async def _fake_get_batch_status_progress(batch_id):
    _progress_poll_count["n"] += 1
    if _progress_poll_count["n"] == 1:
        # still running: one request finished, none failed yet
        return {"processing_status": "in_progress", "request_counts": {
            "processing": 0, "succeeded": 1, "errored": 0, "canceled": 0, "expired": 0,
        }}
    return {"processing_status": "ended", "results_url": "fake://results"}


excel_multi_mod.get_batch_status = _fake_get_batch_status_progress

with open(sample_path, "rb") as f:
    r = check("second large multi-check submits as a batch (for progress test)", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": ""},
    ))
progress_multi_data = r.json()
assert progress_multi_data["status"] == "processing", progress_multi_data
progress_multi_check_id = progress_multi_data["multi_check_id"]

r = check("polling while still in-progress reports live counts instead of finalizing", client.get(
    f"/projects/{project_id}/multi-check/{progress_multi_check_id}", params={"manager_id": regular_id}
))
mid_poll = r.json()
assert mid_poll["status"] == "processing", mid_poll
assert mid_poll["progress"] == {"done": 1, "total": 1}, mid_poll
assert mid_poll["created_at"], mid_poll

r = check("next poll finalizes once Anthropic marks the batch ended", client.get(
    f"/projects/{project_id}/multi-check/{progress_multi_check_id}", params={"manager_id": regular_id}
))
assert r.json()["status"] == "completed", r.json()

excel_multi_mod.get_batch_status = _fake_get_batch_status

# --- cancelling a still-processing upload (Александр asked whether a check
# can be stopped mid-way — e.g. it's taking longer than expected and he'd
# rather re-upload with "Срочно", or he simply changed his mind). Deleting a
# still-processing entry now doubles as "cancel": Anthropic is told to stop
# working on it (so it isn't billed for whatever hadn't started yet) before
# our own record is dropped. ---
_cancel_calls_seen = []


async def _fake_cancel_message_batch(batch_id):
    _cancel_calls_seen.append(batch_id)
    return {"processing_status": "canceling"}


excel_multi_mod.cancel_message_batch = _fake_cancel_message_batch

with open(sample_path, "rb") as f:
    r = check("third large multi-check submits as a batch (for cancel test)", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": ""},
    ))
cancel_multi_data = r.json()
assert cancel_multi_data["status"] == "processing", cancel_multi_data
cancel_multi_check_id = cancel_multi_data["multi_check_id"]

check("cancelling (deleting) a still-processing upload tells Anthropic to stop", client.delete(
    f"/projects/{project_id}/multi-check/{cancel_multi_check_id}", params={"manager_id": regular_id}
))
assert _cancel_calls_seen == ["msgbatch_test123"], _cancel_calls_seen
r = check("cancelled upload no longer in history", client.get(
    f"/projects/{project_id}/multi-check", params={"manager_id": regular_id}
))
assert all(h["id"] != cancel_multi_check_id for h in r.json()), r.json()

# --- rough ETA for a still-processing batch job, learned from how long
# past jobs of a similar size actually took (Александр asked for some kind
# of estimate instead of only elapsed time). A real test run finishes in
# milliseconds, far too fast to naturally produce any meaningful "minutes
# elapsed" — so two fake historical completed batch jobs are inserted
# directly via the DB with known size and duration, and the estimate for a
# new job is checked against the rate they imply. ---
from app.main import _estimate_batch_minutes
from app import models
import datetime as _dt

_hist_db = dbmod.SessionLocal()
try:
    _now_utc = _dt.datetime.now(_dt.timezone.utc)
    _hist_db.add(models.MultiCheck(
        project_id=project_id, filename="hist1.xlsx", source_lang="en",
        status="completed", batch_id="msgbatch_hist1", manager_id=regular_id,
        batch_volume_chars=10_000,
        created_at=_now_utc - _dt.timedelta(minutes=10), completed_at=_now_utc,
    ))
    _hist_db.add(models.MultiCheck(
        project_id=project_id, filename="hist2.xlsx", source_lang="en",
        status="completed", batch_id="msgbatch_hist2", manager_id=regular_id,
        batch_volume_chars=20_000,
        created_at=_now_utc - _dt.timedelta(minutes=20), completed_at=_now_utc,
    ))
    _hist_db.commit()
    # Pooled rate: (10,000 + 20,000) chars over (10 + 20) minutes = 1,000
    # chars/minute — NOT the average of the two jobs' own ratios (which
    # would also happen to be 1,000/min here; the point of pooling is that a
    # small, fast outlier can't dominate the average the way it would if
    # every job's ratio counted equally regardless of size).
    assert _estimate_batch_minutes(_hist_db, 5_000) == 5, _estimate_batch_minutes(_hist_db, 5_000)
    assert _estimate_batch_minutes(_hist_db, 0) is None
finally:
    _hist_db.close()
print("[OK] _estimate_batch_minutes: pools characters/minutes across recent finished batch jobs into "
      "one rate to estimate a new job's duration, rather than averaging each job's own ratio")

with open(sample_path, "rb") as f:
    r = check("multi-check submission now includes a rough ETA once there's history to learn from", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": ""},
    ))
eta_multi_data = r.json()
assert eta_multi_data["status"] == "processing", eta_multi_data
assert eta_multi_data["estimated_minutes"] is not None and eta_multi_data["estimated_minutes"] > 0, eta_multi_data
print("   estimated_minutes:", eta_multi_data["estimated_minutes"])
check("cancel that ETA-test upload so it doesn't linger in history for later tests", client.delete(
    f"/projects/{project_id}/multi-check/{eta_multi_data['multi_check_id']}", params={"manager_id": regular_id}
))

# --- "Срочно" (urgent) flag forces the live/synchronous path even for a
# file whose volume would otherwise route it to the batch queue. The batch
# network calls are still monkeypatched from above, but the urgent path
# must never call them — it should complete in the same request instead of
# coming back as "processing". ---
_batch_calls_seen = []
_original_fake_create_batch = excel_multi_mod.create_message_batch


async def _fake_create_message_batch_tracking(requests):
    _batch_calls_seen.append(requests)
    return await _original_fake_create_batch(requests)


excel_multi_mod.create_message_batch = _fake_create_message_batch_tracking

with open(sample_path, "rb") as f:
    r = check("urgent multi-check bypasses the batch queue", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={
            "source_lang": "",
            "manager_name": "Мария",
            "manager_id": regular_id,
            "extra_instructions": "",
            "urgent": "true",
        },
    ))
urgent_multi_data = r.json()
assert urgent_multi_data["status"] == "completed", urgent_multi_data
assert not _batch_calls_seen, "urgent=true must not go through the batch queue at all"
excel_multi_mod.create_message_batch = _original_fake_create_batch

# --- Александр hit a real production bug: a file with a language column
# header like "fr-CI" but typed with a Cyrillic «с» (U+0441) instead of the
# visually-identical Latin "c" — looks completely normal to a human, but
# Anthropic's Batches API requires custom_id to match ^[a-zA-Z0-9_-]{1,64}$,
# and the old code built custom_id straight from the language string
# ("s{sheet}-{lang}"), so this one bad column got the ENTIRE batch (every
# language in it) rejected by Anthropic with a 400 — which, because it was
# an unhandled exception, came back to the browser as a raw 500 with no
# CORS headers, which Chrome then reported as "blocked by CORS policy",
# completely hiding the real cause. Fixed by building custom_id from the
# sheet/position index only (see build_batch_plan) — never from the
# language string. This proves that holds for ANY weird character, not
# just this one. ---
import re as _re
from app.excel_multi import build_batch_plan as _build_batch_plan_direct
from app.excel_multi import parse_workbook as _parse_workbook_direct

_weird_lang_wb = openpyxl.Workbook()
_weird_lang_ws = _weird_lang_wb.active
_weird_lang_ws.append(["EN", "RU", "fr-сi", "Ünïçø∂€ 漢字"])  # 3rd header: Cyrillic с, not Latin c
_weird_lang_ws.append(["Hello.", "Привет.", "Bonjour.", "Bonjour."])
_weird_lang_buf = io.BytesIO()
_weird_lang_wb.save(_weird_lang_buf)
_weird_lang_buf.seek(0)
_weird_sheets = _parse_workbook_direct(_weird_lang_buf.read())
_weird_requests, _weird_skeleton = _build_batch_plan_direct(_weird_sheets, "en", ["typo"], "", {}, None)
_custom_id_pattern = _re.compile(r"^[a-zA-Z0-9_-]{1,64}$")
_bad_ids = [r["custom_id"] for r in _weird_requests if not _custom_id_pattern.match(r["custom_id"])]
assert not _bad_ids, f"custom_id must always be Anthropic-safe, regardless of the file's own language codes: {_bad_ids}"
assert len(_weird_requests) >= 2, "expected a request for each non-source language, weird characters included"
print(f"[OK] build_batch_plan: custom_id stays ASCII-safe even for language codes with lookalike/unicode characters (e.g. Cyrillic «с» instead of Latin \"c\"): {[r['custom_id'] for r in _weird_requests]}")

# --- and the defensive side of the same fix: an app-wide handler for
# httpx.HTTPError (registered in main.py, not the bare Exception class —
# see its comment for why that distinction is what actually makes CORS
# headers survive) means ANY Anthropic/network failure, anywhere AI checks
# are called, surfaces as a clean, readable error — never a raw crash that
# strips CORS headers and shows up in the browser as a confusing "blocked
# by CORS policy" message with no indication anything is actually wrong
# server-side. Covers all three places that call out to Anthropic:
# batch submission (the exact path Александр's real report hit), the
# live/synchronous multi-check path, and the standalone single-check
# endpoint. ---
import httpx as _httpx_for_fault_injection


async def _fake_create_message_batch_failing(requests):
    raise _httpx_for_fault_injection.ConnectError("simulated Anthropic outage")


_previous_create_batch = excel_multi_mod.create_message_batch
excel_multi_mod.create_message_batch = _fake_create_message_batch_failing
with open(sample_path, "rb") as f:
    r = check("a failed batch submission surfaces as a clean error, not a raw crash", client.post(
        f"/projects/{project_id}/multi-check",
        files={"file": ("Promo_Rules_Localization.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"source_lang": "", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": "", "checks": "typo"},
    ), expect=502)
assert "Не удалось связаться" in r.json()["detail"], r.json()
excel_multi_mod.create_message_batch = _previous_create_batch

# same fault, but on the live/synchronous path (small file, stays under
# BATCH_THRESHOLD_CHARS) — a different code path (run_ai_checks_batch ->
# _call_claude) than the batch-submission one above, so it needs its own
# coverage to actually prove the app-wide handler, not just this one
# call site.
import app.claude_client as claude_client_mod

_previous_call_claude = claude_client_mod._call_claude


async def _fake_call_claude_failing(prompt, model=None):
    raise _httpx_for_fault_injection.ConnectError("simulated Anthropic outage")


claude_client_mod._call_claude = _fake_call_claude_failing
small_wb = openpyxl.Workbook()
small_ws = small_wb.active
small_ws.append(["EN", "RU"])
small_ws.append(["Hello.", "Привет."])
small_buf = io.BytesIO()
small_wb.save(small_buf)
small_buf.seek(0)
r = check("an Anthropic failure on the LIVE multi-check path also surfaces cleanly", client.post(
    f"/projects/{project_id}/multi-check",
    files={"file": ("small.xlsx", small_buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"source_lang": "en", "manager_name": "Мария", "manager_id": regular_id, "extra_instructions": "", "checks": "typo"},
), expect=502)
assert "Не удалось связаться" in r.json()["detail"], r.json()

# and on the standalone single-check endpoint
r = check("an Anthropic failure on the standalone /check endpoint also surfaces cleanly", client.post(
    "/check", json={"source": "Hello.", "translation": "Привет.", "checks": ["typo"]},
), expect=502)
assert "Не удалось связаться" in r.json()["detail"], r.json()
claude_client_mod._call_claude = _previous_call_claude

# --- and the flip side, deliberately: a genuinely unrelated bug (not an
# Anthropic/network failure) must NOT be silently swallowed by the same
# handler — it should still surface as the framework's normal bare 500, so
# a real bug still gets noticed and fixed rather than hidden behind a
# friendly "temporary outage" message forever. Locks in that the handler
# above is registered for httpx.HTTPError specifically, not bare
# Exception. ---
async def _fake_call_claude_unrelated_bug(prompt, model=None):
    raise RuntimeError("some unrelated real bug, not an Anthropic/network failure")


claude_client_mod._call_claude = _fake_call_claude_unrelated_bug
# TestClient's default client re-raises unhandled server exceptions into the
# test process instead of returning them as a response (so a real crash is
# loud in normal testing) — exactly the opposite of what we want to check
# here, so this one call uses its own client with that behaviour turned off.
_no_raise_client = TestClient(app, raise_server_exceptions=False)
r_bug = check(
    "an unrelated bug (not httpx.HTTPError) is NOT masked as a friendly outage message",
    _no_raise_client.post("/check", json={"source": "Hello.", "translation": "Привет.", "checks": ["typo"]}),
    expect=500,
)
assert "Не удалось связаться" not in r_bug.text, r_bug.text
claude_client_mod._call_claude = _previous_call_claude

# --- get_batch_results must not let one garbled .jsonl line (a cut-off
# download, a proxy hiccup) take down parsing of the rest of the batch's
# results — it should skip just that line and keep going. ---
class _FakeBatchResultsResponse:
    def __init__(self, text):
        self.text = text

    def raise_for_status(self):
        pass


class _FakeBatchResultsClient:
    def __init__(self, *args, **kwargs):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    async def get(self, url, headers=None):
        good_line = (
            '{"custom_id": "s0-t0", "result": {"type": "succeeded", '
            '"message": {"content": [{"type": "text", "text": "[]"}], '
            '"usage": {}, "stop_reason": "end_turn"}}}'
        )
        return _FakeBatchResultsResponse("\n".join([good_line, "{not valid json", good_line.replace("t0", "t1")]))


_previous_async_client = claude_client_mod.httpx.AsyncClient
claude_client_mod.httpx.AsyncClient = _FakeBatchResultsClient
_results = asyncio.get_event_loop().run_until_complete(claude_client_mod.get_batch_results("fake://results"))
claude_client_mod.httpx.AsyncClient = _previous_async_client
assert set(_results.keys()) == {"s0-t0", "s0-t1"}, _results
print("[OK] get_batch_results skips a malformed .jsonl line instead of crashing the whole batch")

# --- a bad/expired API key (401/403 from Anthropic) is a config problem on
# our side, not a transient outage — the handler should say so distinctly
# rather than telling the user to just try again in a minute. ---
class _FakeAuthErrorResponse:
    status_code = 401


async def _fake_call_claude_auth_error(prompt, model=None):
    raise _httpx_for_fault_injection.HTTPStatusError(
        "401 Unauthorized", request=None, response=_FakeAuthErrorResponse()
    )


claude_client_mod._call_claude = _fake_call_claude_auth_error
r_auth = check(
    "a 401 from Anthropic (bad/expired API key) gets a distinct 'not a transient outage' message",
    client.post("/check", json={"source": "Hello.", "translation": "Привет.", "checks": ["typo"]}),
    expect=502,
)
assert "ошибк" in r_auth.json()["detail"].lower() and "автор" in r_auth.json()["detail"].lower(), r_auth.json()
assert "временный сбой" not in r_auth.json()["detail"], r_auth.json()
claude_client_mod._call_claude = _previous_call_claude

# --- re-uploading the tone doc replaces it, doesn't accumulate ---
twb3 = openpyxl.Workbook()
tws3 = twb3.active
tws3.append(["EN", "RU"])
tws3.append(["Формальное", "Формальное"])
tbuf3 = io.BytesIO()
twb3.save(tbuf3)
tbuf3.seek(0)
r = check("admin re-uploads tone doc (replaces)", client.post(
    f"/projects/{project_id}/tone/upload",
    files={"file": ("tone2.xlsx", tbuf3, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"manager_id": admin_id},
))
assert r.json()["rule_count"] == 2, r.json()

# --- project template copy: new project starts with the same docs, fully
# independent afterward (editing one never touches the other) ---
r = check("create project copying requirements from the first", client.post(
    "/projects", json={"name": "LS Promo", "manager_id": admin_id, "copy_from_project_id": project_id}
))
copy_project_id = r.json()["id"]
assert r.json()["tone_filename"] == "tone2.xlsx", r.json()

r = check("copied project's tone status matches source", client.get(f"/projects/{copy_project_id}/tone/status"))
assert r.json()["rule_count"] == 2, r.json()

# --- the copy must ALSO inherit the source project's language catalog —
# without this, a project created "from" a template would start with a
# usable tone doc but an empty, useless checkbox list ---
r = check("copied project's language catalog also matches source", client.get(f"/projects/{copy_project_id}/known-languages"))
assert set(r.json()["languages"]) == {"ru", "es-mx", "en", "ar", "kk", "pt-br"}, r.json()
# and it's a genuinely independent copy — removing a language from the
# COPY must not touch the original
check("removing a language from the copy", client.delete(
    f"/projects/{copy_project_id}/languages/kk", params={"manager_id": admin_id}
))
r = check("original project's catalog is untouched by the copy's edit", client.get(f"/projects/{project_id}/known-languages"))
assert "kk" in r.json()["languages"], r.json()

# re-uploading the ORIGINAL project's tone doc must not affect the copy
twb4 = openpyxl.Workbook()
tws4 = twb4.active
tws4.append(["EN", "RU", "ES (MX)"])
tws4.append(["Формальное", "Формальное", "Неформальное"])
tbuf4 = io.BytesIO()
twb4.save(tbuf4)
tbuf4.seek(0)
check("re-upload original project's tone doc again", client.post(
    f"/projects/{project_id}/tone/upload",
    files={"file": ("tone3.xlsx", tbuf4, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    data={"manager_id": admin_id},
))
r = check("copy's tone doc is untouched by the original's re-upload", client.get(f"/projects/{copy_project_id}/tone/status"))
assert r.json()["rule_count"] == 2, r.json()

# --- project deletion requires the admin's password ---
check("delete project wrong password rejected", client.request(
    "DELETE", f"/projects/{copy_project_id}", json={"manager_id": admin_id, "code": "wrong"}
), expect=401)
check("delete project by non-admin rejected", client.request(
    "DELETE", f"/projects/{copy_project_id}", json={"manager_id": regular_id, "code": "9999"}
), expect=403)
check("delete project with correct password", client.request(
    "DELETE", f"/projects/{copy_project_id}", json={"manager_id": admin_id, "code": "1234"}
))
r = check("deleted project no longer listed", client.get("/projects"))
assert len(r.json()) == 1, r.json()

# --- re-running migrations on an already-migrated DB should be a no-op ---
dbmod.init_db()
r = check("managers survive re-migration", client.get("/managers"))
assert len(r.json()) == 2
r = check("projects survive re-migration", client.get("/projects"))
assert len(r.json()) == 1
r = check("tone doc survives re-migration", client.get(f"/projects/{project_id}/tone/status"))
assert r.json()["rule_count"] == 3

# --- unit tests: language-code granularity bridging (resolve_lang_code /
# merge_lang_codes) ---
from app.excel_multi import resolve_lang_code, merge_lang_codes
from app.project_docs import parse_tone_workbook

# exact match wins even when a base-subtag match would also be possible
assert resolve_lang_code("ko-kr", ["ko-kr", "ko"]) == "ko-kr"
# plain code resolves to the one region-qualified entry that shares its base
assert resolve_lang_code("ko", ["ko-KR", "en-us"]) == "ko-KR"
# ...and the reverse direction: a region-qualified request resolves to a
# plain entry sharing its base
assert resolve_lang_code("ko-kr", ["ko", "en-us"]) == "ko"
# several distinct regions for the same base -> refuses to guess
assert resolve_lang_code("es", ["es-ES", "es-AR", "es-MX"]) is None
# unknown language entirely -> no match
assert resolve_lang_code("de", ["ko-KR", "es-MX"]) is None
# several same-base candidates, but with `values` supplied: if they all
# carry the identical rule anyway (Argentina and "rest of Latin America"
# sometimes do), it's safe to resolve rather than refuse — this is the
# real case Александр described: es-MX/es-CL/es-PE-style codes that all
# mean the same "not Spain, not Argentina" format
latam_values = {"es-ar": {"decimal": "—.50"}, "es-mx": {"decimal": "—.50"}}
assert resolve_lang_code("es", ["es-ar", "es-mx"], values=latam_values) == "es-ar"  # identical values -> safe to resolve either way
distinct_values = {"es-es": {"decimal": "—,50"}, "es-ar": {"decimal": "—.50"}}
assert resolve_lang_code("es", ["es-es", "es-ar"], values=distinct_values) is None  # genuinely different -> still refuses
# country-code-style label (Tone's real doc names Kazakh "KZ", Bengali
# "BD", Tajik "TJ" — the COUNTRY, not the ISO language subtag) bridges to
# the real language-region code via the region half, not the base half
assert resolve_lang_code("kz", ["kk-KR", "kk-KZ", "en-US"]) == "kk-KZ"
assert resolve_lang_code("bd", ["bn-BD", "hi-IN"]) == "bn-BD"
print("[OK] resolve_lang_code: exact match, any-subtag fallback (base or region, both "
      "directions), ambiguous multi-region refusal, unknown language, value-equality fallback")

assert merge_lang_codes(["ko", "ko-KR"]) == ["ko-KR"]
assert merge_lang_codes(["es", "es-ES", "es-AR", "es-MX"]) == ["es-AR", "es-ES", "es-MX"]
assert merge_lang_codes(["ru", "fr"]) == ["fr", "ru"]
# the same country-code-style bridging as above, but for the display list
assert merge_lang_codes(["kz", "kk-KZ"]) == ["kk-KZ"]
print("[OK] merge_lang_codes: collapses same-granularity duplicates, "
      "keeps genuinely distinct regional variants, bridges country-code-style "
      "labels to their region subtag, leaves unrelated codes alone")

# --- Александр hit this live: Tone.xlsx has both a bare "AR" (Arabic)
# column and an "ES (AR)" (Spanish, Argentina) column. Arabic's own bare
# code purely coincidentally spells the same two letters as Argentina's
# ISO-3166 country code, which sits inside "es-ar" as its region half —
# the OLD any-subtag matching treated that coincidence as "same language",
# silently merged bare "ar" into the Spanish group, and Arabic vanished
# from the known-languages list entirely. Same landmine, still latent,
# for five more real languages whose code coincidentally spells another
# country's real ISO-3166 code: Bengali/Brunei, Kyrgyz/Cayman Islands,
# Marathi/Mauritania, Tajik/Togo, Tagalog/Timor-Leste. ---
assert "ar" in merge_lang_codes(["ar", "es-ar", "es-mx"])
assert "es-ar" in merge_lang_codes(["ar", "es-ar", "es-mx"])
assert resolve_lang_code("ar", ["es-ar", "es-mx"]) is None  # Arabic must not resolve to a Spanish-Argentina row
for lang, other_country_code in [("bn", "bn"), ("ky", "ky"), ("mr", "mr"), ("tg", "tg"), ("tl", "tl")]:
    merged = merge_lang_codes([lang, f"fr-{other_country_code}", "fr-fr"])
    assert lang in merged, (lang, merged)
    assert f"fr-{other_country_code}" in merged, (lang, merged)
# but a bare code that ISN'T itself a real language (the agency's own
# country-code-style shorthand, same as the KZ/BD tests above) still
# bridges via region exactly as before — this fix must not have broken
# the legitimate case it was carved out of.
assert merge_lang_codes(["kg", "ky-KG"]) == ["ky-KG"]
assert resolve_lang_code("kg", ["ky-KG", "en-US"]) == "ky-KG"
# and a bare code that genuinely IS the same language as a region-qualified
# one (not a coincidence) still merges normally.
assert merge_lang_codes(["ar", "ar-eg"]) == ["ar-eg"]
# same coincidence, a sixth real case: Александр's files use bare "my" for
# Malay specifically (not ISO-639's own Burmese meaning — see
# claude_client.LANG_CODE_MEANING_OVERRIDES), which is itself also
# Malaysia's real ISO-3166 country code — bare "my" must bridge to a
# genuine "my-*" variant of Malay, but never get pulled into an unrelated
# language's Malaysia-region entry (e.g. "zh-MY") just because they
# happen to share those two letters.
assert merge_lang_codes(["my", "zh-MY", "zh-CN"]) == ["my", "zh-CN", "zh-MY"]
assert resolve_lang_code("my", ["zh-MY", "zh-CN"]) is None
print("[OK] merge_lang_codes/resolve_lang_code: a real independent language "
      "whose code coincidentally spells another country's ISO-3166 code "
      "(Arabic \"ar\" vs Argentina inside \"es-ar\", and the same latent risk "
      "for Bengali, Kyrgyz, Marathi, Tajik, Tagalog) is never absorbed into "
      "or resolved against an unrelated language's region — while the "
      "legitimate country-code-shorthand bridging (KZ/BD/KG-style) and "
      "genuine same-language bridging (ar/ar-eg) both still work")

# --- _lang_selected: decides whether a manager's target_langs_filter (raw
# catalog-checkbox codes) covers a given FILE column — used by the actual
# check run, unlike resolve_lang_code's own base-language shortcut (which
# is fine for merge_lang_codes' display grouping) it must NEVER bridge two
# codes that are BOTH already region-qualified just because they share a
# base language: "es-mx" ticked must never silently also check an "es-es"
# file column, since those are deliberately distinct, explicitly-added
# catalog languages. It still bridges a bare code against a region-
# qualified one either direction (a ticked bare "pt" covers a file's
# "pt-br" column, and vice versa), including the country-code-shorthand
# case (ticked "kz" covers a file's "kk-KZ" column) — exactly the
# Portuguese-default scenario this was added for — and refuses instead of
# guessing when more than one filter entry would bridge. ---
from app.excel_multi import _lang_selected

assert _lang_selected("ru", {"ru", "es-mx"}) is True  # exact match
assert _lang_selected("es-mx", {"ru", "es-mx"}) is True  # exact match
assert _lang_selected("pt-br", {"pt"}) is True  # bare filter covers a region-qualified file column
assert _lang_selected("pt", {"pt-br"}) is True  # and the reverse direction
assert _lang_selected("kk-KZ", {"kz"}) is True  # country-code-shorthand bridging, same as resolve_lang_code
assert _lang_selected("es-es", {"ru", "es-mx"}) is False  # NEVER bridge two distinct explicit regions
assert _lang_selected("ar", {"es-ar"}) is False  # Arabic must never match Argentina's region code
assert _lang_selected("es", {"es-ar", "es-mx"}) is False  # ambiguous bridge -> refuse, don't guess
# ...but NOT the other way around, deliberately: a bare catalog tick ("es")
# against a file with BOTH "es-ar" and "es-mx" columns selects each one
# independently rather than refusing both — over-including a language is a
# far smaller problem than the silent-drop bug this function exists to fix.
assert _lang_selected("es-ar", {"es"}) is True
assert _lang_selected("es-mx", {"es"}) is True
print("[OK] _lang_selected: bridges a bare catalog code against a file's region-qualified "
      "column (and vice versa), same as resolve_lang_code's bare/region bridging, but never "
      "bridges two codes that are both already region-qualified, refuses when a file column "
      "could ambiguously bridge to several ticked filter entries, and (deliberately, "
      "asymmetrically) still selects every explicit file column that bridges to one bare "
      "catalog tick rather than refusing to guess in that direction")

# --- Александр's Portuguese is always Brazilian, never Portugal's — a bare
# "PT" column (or a manually-typed catalog addition of just "pt") must
# default to "pt-br", the same way "ES (AR)" already tells the AI check it's
# Argentine Spanish rather than leaving a bare "es" to be guessed at. An
# EXPLICIT region must never be overridden by that default, whichever style
# it's written in. ---
from app.excel_multi import _normalize_lang_label

assert _normalize_lang_label("PT") == "pt-br"
assert _normalize_lang_label("pt") == "pt-br"
assert _normalize_lang_label("PT (PT)") == "pt-pt"  # explicit region always wins over the default
assert _normalize_lang_label("PT-BR") == "pt-br"  # already-explicit form passes through unchanged
assert _normalize_lang_label("PT (BR)") == "pt-br"
assert _normalize_lang_label("ES") == "es"  # the default is Portuguese-specific, not applied to other languages
print("[OK] _normalize_lang_label: a bare \"pt\"/\"PT\" defaults to Brazilian Portuguese "
      "(\"pt-br\"), while any explicitly-written region is always left alone")

# --- pick_source_lang must bridge the same granularity mismatches as
# everything else in this file (via resolve_lang_code), not do a literal
# string match — Александр hit this live: he picked "Русский" as the
# source language, but his file's Russian column wasn't spelled exactly
# "ru" (a region-qualified "ru-RU"), the old literal check silently missed
# it, and the source language silently fell back to whatever column was
# labeled "en-001" instead — his check ran source-vs-source against the
# wrong pair of columns without any error or warning. ---
from app.excel_multi import pick_source_lang

fake_sheets = [{"languages": ["ru-RU", "en-001", "kz"]}]
assert pick_source_lang(fake_sheets, "ru") == "ru-RU", pick_source_lang(fake_sheets, "ru")
assert pick_source_lang(fake_sheets, "kk-KZ") == "kz", pick_source_lang(fake_sheets, "kk-KZ")
# still falls back to English, then alphabetically first, when the
# requested language genuinely isn't in the file at all
assert pick_source_lang(fake_sheets, "de") == "en-001", pick_source_lang(fake_sheets, "de")
assert pick_source_lang([{"languages": ["kz", "en-001"]}], None) == "en-001"
print("[OK] pick_source_lang: resolves the manager's chosen source language against "
      "a differently-granular spelling of the same language in the file, instead of "
      "silently falling back to English")

# parse_tone_workbook: real layout is column-per-language, not
# row-per-language as originally assumed — including a country-code
# label ("KZ") and a combined "/"-separated header applying to two codes
twb2 = openpyxl.Workbook()
tws2 = twb2.active
tws2.append(["EN", "RU", "KZ", "FR-CI / FR-FR", "ES (MX)"])
tws2.append(["Формальное", "Формальное", "Формальное", "Неформальное", "Неформальное"])
tbuf2 = io.BytesIO()
twb2.save(tbuf2)
tbuf2.seek(0)
tone_rows = {r["lang_code"]: r["register"] for r in parse_tone_workbook(tbuf2.read())}
assert tone_rows == {
    "en": "formal", "ru": "formal", "kz": "formal",
    "fr-ci": "informal", "fr-fr": "informal", "es-mx": "informal",
}, tone_rows
print("[OK] parse_tone_workbook: column-per-language layout (not row-per-language), "
      "combined header split across both codes")

# --- model tiering: confirmed "hard" languages get the stronger model,
# matched by base subtag so any region variant of them qualifies too ---
from app.claude_client import _model_for_lang
from app.config import settings

for hard in ["kk", "kk-KZ", "ky-KG", "tg-TJ", "uz", "sw-KE", "te-IN", "mr-IN", "az-AZ"]:
    assert _model_for_lang(hard) == settings.CLAUDE_MODEL_HARD, hard
for normal in ["ru", "es-mx", "en", "de-DE", "fr"]:
    assert _model_for_lang(normal) == settings.CLAUDE_MODEL, normal
print("[OK] _model_for_lang: confirmed hard-language list (kk/ky/tg/uz/sw/te/mr/az) "
      "routes to CLAUDE_MODEL_HARD by base subtag, everything else to CLAUDE_MODEL")

# --- "my" is ISO-639's code for Burmese, but Александр's files use it for
# Malay — left unclarified, the model assumes Burmese and reports correct
# Malay text as being in the wrong language. The prompt must explicitly
# override this one code's meaning instead of just stating the raw code. ---
from app.claude_client import _target_lang_line

my_line = _target_lang_line("my")
assert "малайск" in my_line.lower(), my_line
assert "бирманск" in my_line.lower(), my_line  # explicitly rules out the ISO-standard meaning
normal_line = _target_lang_line("ru")
assert "малайск" not in normal_line.lower() and "бирманск" not in normal_line.lower(), normal_line
print("[OK] _target_lang_line: the «my» code is explicitly clarified as Malay (not the ISO-standard "
      "Burmese) so the model doesn't misjudge correct Malay text as the wrong language")

# --- numbers check: a correctly localized decimal comma or zero-padded
# hour must NOT be flagged as a mismatch — Александр hit this live: an
# Azerbaijani translation writing "0,40" for the source's "$0.40" and
# "03:00" for the source's "3:00" was flagged "numbers" even though
# nothing was actually mistranslated, just correctly reformatted. ---
from app.rule_checks import check_numbers

az_source = (
    "To participate in the tournament, you must confirm participation in any game from the list, "
    "place bets from $0.40, and complete missions. The tournament runs from 09/22/2026 (3:00 UTC) "
    "to 09/28/2026 (22:59 UTC). The full rules are available in the in-game menu. Minimum bet: $0.40"
)
az_translation = (
    "Turnirdə iştirak etmək üçün siyahıdakı istənilən oyunda iştirakınızı təsdiqləməli, 0,40 ₼ və "
    "daha çox mərc etməli və missiyaları yerinə yetirməlisiniz. Turnir from 22/09/2026 (03:00 UTC) – "
    "28/09/2026 (22:59 UTC) tarixləri arasında keçirilir. Tam qaydaları oyun içi menyuda oxuya "
    "bilərsiniz. Min. mərc: 0,40 ₼"
)
assert check_numbers(az_source, az_translation) == [], check_numbers(az_source, az_translation)
# A genuine mismatch (thousands grouping aside — see below) must still fire.
assert check_numbers("The bonus is $50.", "Бонус составляет $500.") != []
# A comma used to GROUP thousands (not as a decimal separator) is left
# alone — "50,000" really is a different number from "50" if unmatched.
assert check_numbers("$50,000 prize", "50 тысяч приз") != []
print("[OK] check_numbers: decimal-comma and zero-padded-hour localization no longer "
      "false-flagged as a numbers mismatch; real mismatches and thousands-grouping still caught")

# --- dropping (or keeping) a thousands-grouping comma must not be flagged
# either — Александр hit this live: a translation correctly kept some of a
# promo's big numbers grouped ("1,500,000") but wrote a smaller one
# ungrouped ("1400" for the source's "1,400"), and it was flagged as a
# mismatch even though the value never changed. ---
assert check_numbers("Win up to $1,400 today", "Выиграйте до $1400 сегодня") == []
assert check_numbers("Prize: $1,500,000", "Приз: $1500000") == []
# But an actually different grouped number must still be caught.
assert check_numbers("Win up to $1,400 today", "Выиграйте до $1,500 сегодня") != []
print("[OK] check_numbers: a thousands-grouping comma can be freely added or dropped "
      "without being flagged, while an actually different grouped number is still caught")

# --- a SPACE is just as legitimate a thousands separator as a comma —
# it's how Russian formats a big number ("1 500 000") — but NUMBER_RE
# can't include a bare space (that would merge unrelated numbers separated
# by ordinary whitespace). Александр hit this live in a real prize-table
# promo: every single big number came back "different" between Russian
# ("1 500 000") and English/Spanish ("1,500,000") for no reason at all. ---
assert check_numbers("1st place — 1,500,000 ARS", "1 место — 1 500 000 ARS") == []
assert check_numbers("Prize: 90,000 ARS", "Приз: 90 000 ARS") == []
# An actually different space-grouped number must still be caught, and the
# message must point at the specific numbers that differ, not dump every
# number in the text — a long paragraph can have 20+ numbers where only
# one is actually wrong.
mismatch = check_numbers("70th-99th place — 70,000 ARS", "С 70 по 99 место — 72 000 ARS")
assert mismatch, mismatch
assert "70000" in mismatch[0]["message"] and "72000" in mismatch[0]["message"], mismatch
print("[OK] check_numbers: a space used to group thousands (standard Russian formatting) "
      "is treated the same as a comma — freely interchangeable without being flagged — "
      "while an actually different number is still caught, and the message names the "
      "specific number(s) that differ rather than dumping the whole list")

# --- a PERIOD is just as legitimate a thousands separator as a comma or a
# space — German, Spanish and several other target languages group
# thousands that way ("50.000", "1.500.000"). Александр hit this live: the
# report showed source "50000" and translation "50.000" as two different
# numbers even though the value never changed, purely because the comma
# rule above only covered the comma. Also covers a currency symbol whose
# position and spacing changes between source and translation ("$100000"
# vs "100.000$") — the symbol itself was never part of the extracted
# number to begin with, so its placement was already a non-issue; this
# just confirms it stays that way once the number itself is also
# period-grouped. ---
assert check_numbers("Win up to 50000 today", "Выиграйте до 50.000 сегодня") == []
assert check_numbers("Prize: $1500000", "Приз: 1.500.000$") == []
assert check_numbers("Min. bet: $100000", "Мин. ставка: 100.000 $") == []
# But an actually different period-grouped number must still be caught.
mismatch_dot = check_numbers("Win up to 50000 today", "Выиграйте до 55.000 сегодня")
assert mismatch_dot, mismatch_dot
assert "50000" in mismatch_dot[0]["message"] and "55000" in mismatch_dot[0]["message"], mismatch_dot
print("[OK] check_numbers: a period used to group thousands (standard German/Spanish "
      "formatting) is treated the same as a comma or a space — freely interchangeable "
      "without being flagged, including when a currency symbol moves position/spacing "
      "along with it — while an actually different number is still caught")

# --- a DATE must not be flagged just because the target language writes
# the day/month/year in a different order and/or with a different
# separator — Александр hit this live: a source date "09/22/2026"
# (MM/DD/YYYY, slash-separated) was correctly localized as "22.09.2026"
# (DD.MM.YYYY, dot-separated), and the check flagged it as a numbers
# mismatch even though every digit was correct — only the grouping/order
# changed, which is expected and fine. ---
date_source = "The tournament runs from 09/22/2026 (3:00 UTC) to 09/28/2026 (22:59 UTC)."
date_translation_ru = "Турнир проходит с 22.09.2026 (03:00 UTC) по 28.09.2026 (22:59 UTC)."
assert check_numbers(date_source, date_translation_ru) == [], check_numbers(date_source, date_translation_ru)
# But an actually WRONG date (a real digit changed, not just reordered)
# must still be caught.
date_translation_wrong = "Турнир проходит с 23.09.2026 (03:00 UTC) по 28.09.2026 (22:59 UTC)."
assert check_numbers(date_source, date_translation_wrong) != []
print("[OK] check_numbers: a date's day/month/year order and separator style "
      "(dots vs slashes) can change freely between languages without being flagged, "
      "while an actual wrong date digit is still caught")

# --- the "typo" AI check catches a WRONG CURRENCY entirely (e.g. euro
# instead of dollar) as a genuine translation error, not just a stylistic
# quirk — Александр hit a real case where $0.40 was mistranslated as
# "0,40 €" ---
from app.claude_client import CHECK_LABELS

assert "ДРУГАЯ ВАЛЮТА" in CHECK_LABELS["typo"], CHECK_LABELS["typo"]
print("[OK] currency identity (wrong currency, e.g. € instead of $) is scoped to «опечатки/ошибки»")

# --- a plain misspelling in the translation itself ("resulits" for
# "results") must be in scope too, even though the meaning is still
# perfectly clear from context — Александр hit this live: the AI didn't
# report it because the old wording scoped "опечатки/ошибки" to ONLY
# meaning-distorting errors, which technically excludes an obvious
# spelling slip a reader can still understand. The label now explicitly
# names plain misspellings as their own always-in-scope category,
# distinct from a stylistic/synonym choice (which stays out of scope). ---
assert "неправильно написанное слово" in CHECK_LABELS["typo"], CHECK_LABELS["typo"]
print("[OK] a plain spelling mistake is explicitly in scope for «опечатки/ошибки» even when the "
      "meaning is still clear from context, distinct from a stylistic/synonym choice which isn't")

# --- but when the free "numbers" rule check is ALSO running in the same
# request, the AI must not also report a plain digit/date mismatch under
# "typo" — that's the exact duplicate Александр hit: the same wrong-year
# date shown once as "numbers" and again, reworded, as "typo". The AI is
# told to leave plain digits to the rule check and only still flag
# currency IDENTITY (not itself a digit) under "typo". When "numbers"
# ISN'T part of this run, the AI keeps acting as the sole backstop for a
# wrong number, exactly as before this fix. ---
from app.claude_client import _calibration

calib_with_numbers = _calibration(["typo", "numbers"])
calib_without_numbers = _calibration(["typo"])
assert "не сообщай о них здесь" in calib_with_numbers, calib_with_numbers
assert "не сообщай о них здесь" not in calib_without_numbers, calib_without_numbers
assert "валюта или число не совпадают" in calib_without_numbers, calib_without_numbers
print("[OK] the AI is told to skip plain digit/date mismatches under «опечатки/ошибки» whenever "
      "the free «numbers» rule check is also part of the same run (avoids the same error being "
      "reported twice under two different labels), and stays the sole backstop for numbers "
      "otherwise")

# --- the prompt also tells the model not to squeeze an out-of-scope
# finding into whichever type happens to be the only one allowed —
# Александр hit exactly this with the (now-removed) glossary check, but
# the instruction itself is generic, not glossary-specific, so it stays
# relevant for any single-check-type run. ---
from app.claude_client import SINGLE_PROMPT, BATCH_PROMPT

assert "Не подгоняй" in SINGLE_PROMPT and "Не подгоняй" in BATCH_PROMPT
print("[OK] the prompt explicitly forbids squeezing an out-of-scope finding into whichever "
      "type happens to be the only one allowed")

# --- AI findings are hard-filtered to only the checks actually requested,
# even if the model ignores the prompt's instruction and reports something
# else anyway ---
from app.claude_client import _allowed_ai_types, _filter_findings_by_checks, run_ai_checks
import app.claude_client as claude_client_mod

assert _allowed_ai_types(["typo"]) == {"typo"}
assert _allowed_ai_types(["typo", "punctuation", "max_length"]) == {"typo"}  # rule checks aren't AI types

raw_findings = [
    {"type": "typo", "severity": "medium", "message": "неверная валюта в переводе"},
    {"type": "untranslatable", "severity": "high", "message": "слово не переведено"},
]
assert _filter_findings_by_checks(raw_findings, ["typo"]) == [raw_findings[0]]

settings.ANTHROPIC_API_KEY = "fake-key-for-smoketest"
captured_prompts = []


async def _fake_call_claude(prompt, model=None):
    captured_prompts.append(prompt)
    # Simulates a model that ignores "проверяй только typo" and
    # reports an untranslatable-text issue anyway.
    text = (
        '[{"type": "typo", "severity": "medium", "message": "неверная валюта в переводе"},'
        '{"type": "untranslatable", "severity": "high", "message": "слово не переведено"}]'
    )
    return (text, {"input_tokens": 500, "output_tokens": 100}, "end_turn")


claude_client_mod._call_claude = _fake_call_claude
findings, ai_cost = asyncio.get_event_loop().run_until_complete(
    run_ai_checks(
        "source", "translation", ["typo"], target_lang="az-az",
    )
)
assert findings == [raw_findings[0]], findings
assert ai_cost > 0, ai_cost
# The JSON schema shown to the model is also scoped down to just the
# requested check(s), not a fixed always-all list.
type_enum_line = next(line for line in captured_prompts[0].splitlines() if '"type":' in line)
assert "typo" in type_enum_line and "untranslatable" not in type_enum_line, type_enum_line
print("[OK] AI findings hard-filtered to requested checks even when the model reports "
      "an out-of-scope finding anyway (prompt's type list is also scoped down, in addition)")

# --- a truncated JSON response (the model hit the max_tokens ceiling
# mid-array) must not lose every finding that came before the cut — only
# the incomplete trailing entry is dropped, everything complete survives —
# and callers must be told the response was truncated at all, rather than
# a partial result silently looking like a complete "checked, nothing
# else" report. Александр hit exactly this: a large Spanish multi-check
# came back with an incomplete report and no indication anything was cut
# short. ---
from app.claude_client import parse_json_array, _truncation_warning, _ai_failure_warning

truncated_text = (
    '[{"type": "typo", "severity": "medium", "message": "первая находка"},'
    '{"type": "typo", "severity": "high", "message": "вторая находка"},'
    '{"type": "typo", "sever'  # cut off mid-object, exactly as a max_tokens cutoff would do
)
salvaged = parse_json_array(truncated_text)
assert len(salvaged) == 2, salvaged
assert salvaged[0]["message"] == "первая находка" and salvaged[1]["message"] == "вторая находка", salvaged
print("[OK] parse_json_array: a JSON array cut off mid-object (max_tokens truncation) "
      "keeps every complete finding before the cut instead of losing the whole response")


async def _fake_call_claude_truncated(prompt, model=None):
    return ('[{"type": "typo", "severity": "medium", "message": "неверная валюта"}]',
            {"input_tokens": 500, "output_tokens": 100}, "max_tokens")


claude_client_mod._call_claude = _fake_call_claude_truncated
findings_trunc, _ = asyncio.get_event_loop().run_until_complete(
    run_ai_checks("source", "translation", ["typo"], target_lang="az-az")
)
assert any(f["type"] == "system" for f in findings_trunc), findings_trunc
claude_client_mod._call_claude = _fake_call_claude
print("[OK] run_ai_checks adds a visible system warning whenever the AI response was "
      "cut off by the max_tokens ceiling, instead of silently showing a partial result")

# --- same guarantee on the async Message Batches merge path
# (finalize_batch_results): a language whose batch response was truncated,
# or whose request errored/expired/never came back at all, gets a visible
# warning finding rather than silently showing as "checked, nothing found".
from app.excel_multi import finalize_batch_results

fake_skeleton = {
    "checks": ["typo"],
    "source_lang": "en",
    "sheets": [{
        "sheet_name": "Sheet1",
        "target_langs": ["es-mx", "fr", "de"],
        "unrecognized_columns": [],
        "row_count": 1,
        "languages": {
            "es-mx": {
                "custom_id": "s0-es-mx", "model": "claude-haiku-4-5-20251001",
                "number_to_index": {}, "rows": [{"excel_row": 2, "context": "", "source": "a", "translation": "b", "findings": []}],
            },
            "fr": {
                "custom_id": "s0-fr", "model": "claude-haiku-4-5-20251001",
                "number_to_index": {}, "rows": [{"excel_row": 2, "context": "", "source": "a", "translation": "b", "findings": []}],
            },
            "de": {
                # never made it into the results at all (e.g. dropped between submit and poll)
                "custom_id": "s0-de", "model": "claude-haiku-4-5-20251001",
                "number_to_index": {}, "rows": [{"excel_row": 2, "context": "", "source": "a", "translation": "b", "findings": []}],
            },
        },
    }],
}
fake_batch_results = {
    "s0-es-mx": {"text": "[", "usage": {"input_tokens": 10, "output_tokens": 8000}, "stop_reason": "max_tokens", "result_type": "succeeded"},
    "s0-fr": {"text": None, "usage": {}, "stop_reason": None, "result_type": "errored"},
    # "s0-de" deliberately absent
}
merged = finalize_batch_results(fake_skeleton, fake_batch_results)
by_lang = merged["sheets"][0]["languages"]
assert any(f["type"] == "system" for row in by_lang["es-mx"] for f in row["findings"]), by_lang["es-mx"]
assert any(f["type"] == "system" for row in by_lang["fr"] for f in row["findings"]), by_lang["fr"]
assert any(f["type"] == "system" for row in by_lang["de"] for f in row["findings"]), by_lang["de"]
print("[OK] finalize_batch_results: a truncated, errored, or missing-entirely batch result "
      "each surface a visible system warning instead of silently reporting as clean")

# --- the /check endpoint itself surfaces the real AI cost, not just the
# internal run_ai_checks helper (Александр asked to see the cost of each
# check after it runs) ---
settings.ANTHROPIC_API_KEY = "fake-key-for-smoketest"
r = check("standalone /check surfaces cost_usd for an AI-backed check", client.post(
    "/check",
    json={
        "source": "source text",
        "translation": "translation text",
        # "typo" needs no project document to run.
        "checks": ["typo"],
    },
))
check_cost_data = r.json()
assert check_cost_data["cost_usd"] > 0, check_cost_data
print(f"   /check cost_usd: {check_cost_data['cost_usd']}")
settings.ANTHROPIC_API_KEY = ""

# --- parse_workbook: obviously-not-a-language columns (per-channel
# character-limit spec columns, "ТЗ", etc.) are dropped silently instead of
# being dumped into the noisy "not recognized as languages" notice —
# Александр flagged a real file whose limit columns ("NOTIF title: 20",
# "Лимиты: PUSH banner: 25", ...) were cluttering that message even though
# it's obvious on sight they're not languages.
from app.excel_multi import parse_workbook
import openpyxl as _openpyxl

wb_limits = _openpyxl.Workbook()
ws_limits = wb_limits.active
ws_limits.title = "Sheet1"
ws_limits.append([
    "Context", "ТЗ", "Лимиты: NOTIF title: 20", "NOTIF banner: 25",
    "NOTIF text: 200", "Лимиты", "en", "ru",
])
ws_limits.append(["greeting", "some brief", "", "", "", "", "Hello", "Привет"])
buf_limits = io.BytesIO()
wb_limits.save(buf_limits)
parsed_limits = parse_workbook(buf_limits.getvalue())
assert len(parsed_limits) == 1, parsed_limits
sheet_limits = parsed_limits[0]
assert sheet_limits["unrecognized_columns"] == [], sheet_limits["unrecognized_columns"]
assert set(sheet_limits["languages"]) == {"en", "ru"}, sheet_limits["languages"]
print("[OK] parse_workbook: per-channel character-limit columns (\"label: number\"), a bare "
      "«Лимиты» header, and «ТЗ» are dropped silently rather than flagged as unrecognized languages")

print("\nALL SMOKETEST CHECKS PASSED")
