"""
Parsing and checking of "Мульти" Excel uploads — the format Crowdin exports:
one sheet (sometimes several) with a header row of language codes, a
"Context" column naming each string, an optional "Max. length" column, and
one row per translatable string.
"""
import asyncio
import io
import re

import openpyxl

from app.claude_client import (
    _ai_failure_warning,
    _filter_findings_by_checks,
    _model_for_lang,
    _truncation_warning,
    _usage_cost,
    build_batch_prompt,
    cancel_message_batch,
    create_message_batch,
    get_batch_results,
    get_batch_status,
    group_batch_findings,
    parse_json_array,
    run_ai_checks_batch,
)
from app.rule_checks import run_rule_checks

LANG_CODE_RE = re.compile(r"^[a-z]{2,3}(-[a-z0-9]{2,5})?$")
AI_CONCURRENCY = 5

# Industry-standard placeholder: a translator (or the client) can mark a
# specific cell "DO NOT TRANSLATE" to mean this string is deliberately left
# as-is for this language on purpose (a brand name, a code, a string that's
# only needed in one of several languages) — not a missing or wrong
# translation. Александр's files use this in exactly that way: some rows
# need translating into every language, others explicitly don't for a given
# one. Recognized case-insensitively, with or without surrounding brackets.
_DO_NOT_TRANSLATE_RE = re.compile(r"^[\[\(]?\s*do\s+not\s+translate\s*[\]\)]?$", re.IGNORECASE)


def _is_do_not_translate(text: str) -> bool:
    return bool(_DO_NOT_TRANSLATE_RE.match((text or "").strip()))

# Above this much combined text (characters, summed across every checkable
# row × every target language — a rough proxy for total AI cost and how
# long a synchronous run would take), a multi-check is submitted through
# Anthropic's Message Batches API instead of run live: half the per-token
# price, but results typically land within an hour instead of a couple of
# minutes. A short single-language check, or a handful of languages, stays
# under this and runs instantly as before — but a small document checked
# across many languages (which is what actually drives the per-token cost
# up) will typically cross this threshold too. This is a starting guess,
# not a measured number — tune it based on real usage once there's a
# track record of actual costs and turnaround times.
BATCH_THRESHOLD_CHARS = 10_000

# Known non-language metadata column names seen in Crowdin-style exports.
# Anything NOT in this list and not Context/Max length is treated as a
# language column — real client files have non-standard codes (a 4-letter
# code, a look-alike Cyrillic character typo'd into a code, etc.) that a
# strict regex would silently drop, which is worse than being permissive.
META_COL_NAMES = {
    "context", "key", "id", "string id", "identifier", "comment",
    "status", "screenshot", "reference", "notes", "note",
    "контекст", "ключ", "тз", "комментарий", "примечание", "статус",
}

# A "label: number" shape — "NOTIF title: 20", "Лимиты: PUSH banner: 25" —
# is unambiguously a character-limit/spec column, never a language: a real
# language code never contains a colon. Matched at the end of the header so
# a "Лимиты: ..." prefix in front doesn't matter. Also treat any header that
# starts with the Russian word for "limits" (a merged section header like a
# bare "Лимиты" spanning several columns, with no colon of its own) the
# same way. Client files pack in columns like this alongside the real
# language columns — Александр flagged them as noise in the "not
# recognized as languages" notice, since it's obvious on sight (and by
# comparing with the source column) that they were never meant to be one.
_LIMIT_SPEC_RE = re.compile(r":\s*\d+\s*$")


def _is_context_col(header: str) -> bool:
    return header.strip().lower() == "context"


def _is_max_length_col(header: str) -> bool:
    h = header.strip().lower()
    return "max" in h and "length" in h


def _is_meta_col(header: str) -> bool:
    return header.strip().lower() in META_COL_NAMES


def _is_limit_spec_col(header: str) -> bool:
    h = header.strip()
    return bool(_LIMIT_SPEC_RE.search(h)) or h.lower().startswith("лимит")


_PAREN_LANG_RE = re.compile(r"^([a-zA-Zа-яА-Я]{2,3})\s*\(([a-zA-Zа-яА-Я0-9]{1,5})\)$")

# The same "ES (MX)"/"PT (BR)" display style, but without the parentheses —
# "ES MX", "PT BR" — a header written this way isn't currently reachable at
# all: with no hyphen and no parens for _PAREN_LANG_RE to key off, it falls
# straight through to a plain .lower() with the space still in it, which
# parse_workbook's own space check then rejects as "looks like prose, not a
# language code". Deliberately requires BOTH words to be fully uppercase in
# the source file (unlike the parenthesized form above, which is
# case-insensitive) — that's what tells a genuine "ES MX"/"PT BR"-style code
# apart from an ordinary two-word column header like "Task name" or a
# capitalized "No data", which are never written in ALL CAPS in Александр's
# real files. Deliberately does NOT try to be clever about the second
# word's length ("PT BR" and "ES MX" happen to be 2+2, but keep the same
# 1-5 character allowance as the parenthesized form for other real cases).
_SPACE_LANG_RE = re.compile(r"^([A-Z]{2,3})\s+([A-Z0-9]{1,5})$")


# Unlike Spanish (which the agency always spells out by region — es-ar,
# es-mx, es-es — because it genuinely handles several), Portuguese has only
# ever meant one thing in Александр's work: Brazilian Portuguese. But
# Portugal's and Brazil's Portuguese differ enough (vocabulary, formality
# conventions) that the AI check should be told explicitly which one it's
# dealing with, exactly the way an explicit "es-ar" already tells it
# Argentine Spanish rather than leaving that to be guessed from a bare "es"
# (see _target_lang_line in app.claude_client — whatever code reaches it as
# target_lang is what the AI is told the language IS). So a bare,
# unqualified "pt" — no region attached at all — defaults to Brazilian
# Portuguese ("pt-br") wherever a language code is first minted: a file
# column header, or a manually-typed catalog addition. An EXPLICIT region
# ("PT (PT)", "pt-pt") is left alone — if Portugal's own Portuguese is ever
# actually needed, spelling it out that way is how to ask for it instead of
# getting the Brazilian default.
_DEFAULT_REGION_FOR_BARE_LANG = {
    "pt": "pt-br",
}


def _normalize_lang_label(label: str) -> str:
    """Turns a display-style language header like "ES (MX)", "PT (BR)", or
    the same without parentheses ("ES MX", "PT BR" — must be ALL CAPS, see
    _SPACE_LANG_RE) into the hyphenated form used everywhere else
    ("es-mx", "pt-br"), while leaving an already-plain code like "es-mx" or
    "fr-сi" untouched. Some client files use one style, some another, so
    all of them need to resolve to the same lang_code. A handful of bare
    codes also get defaulted to a specific region here — see
    _DEFAULT_REGION_FOR_BARE_LANG — since an explicit region ("PT (BR)",
    "pt-pt") above already means something specific and must never be
    overridden by that default."""
    label = label.strip()
    m = _PAREN_LANG_RE.match(label)
    if m:
        return f"{m.group(1).lower()}-{m.group(2).lower()}"
    m = _SPACE_LANG_RE.match(label)
    if m:
        return f"{m.group(1).lower()}-{m.group(2).lower()}"
    code = label.lower()
    return _DEFAULT_REGION_FOR_BARE_LANG.get(code, code)


def _base_lang(code: str) -> str:
    """The base language subtag of a code — "ko" from "ko-KR", "es" from
    "es-mx", or the whole thing if it has no region part."""
    return code.strip().lower().split("-")[0]


def _subtags(code: str) -> set[str]:
    """Every hyphen-separated piece of a code, lowercased — {"kk", "kz"}
    for "kk-KZ". Used to match a bare code against either the language part
    OR the region part of a fuller one, since the agency's own documents
    are inconsistent about which they use for a short label — the Tone doc
    labels some languages by their country's code alone ("KZ" for Kazakh,
    "TJ" for Tajik, "BD" for Bengali) rather than the actual ISO language
    subtag ("kk", "tg", "bn") that a region-qualified code like "kk-KZ"/
    "tg-TJ"/"bn-BD" would use — a plain prefix-only match would miss these
    entirely, since "kz" isn't the base of "kk-kz"."""
    return set(code.strip().lower().split("-"))


# Bare codes that are real, independent ISO-639 languages in this project's
# own right — every one of them actually appears as its own language column
# somewhere in Александр's real files. Used by _language_subtags_compatible
# below to tell apart two very different reasons a short code might share
# letters with a region subtag: the agency's own country-code-style
# shorthand for a language (Tone doc's "KZ" for Kazakh, "TJ" for Tajik,
# "BD" for Bengali — none of which are themselves real ISO-639 codes, so
# matching them against a region subtag is exactly the intended trick), vs
# a genuine ISO-639 language code that PURELY BY COINCIDENCE also spells a
# real but unrelated country's ISO-3166 code (Arabic "ar" vs Argentina's
# country code "AR"; also latent landmines for the same reason even though
# no real file has hit them yet: Bengali "bn"/Brunei "BN", Kyrgyz "ky"/
# Cayman Islands "KY", Marathi "mr"/Mauritania "MR", Tajik "tg"/Togo "TG",
# Tagalog "tl"/Timor-Leste "TL"). A code in this set must never be treated
# as merely someone else's region fragment.
INDEPENDENT_LANGUAGE_CODES = {
    "ar", "az", "bn", "de", "el", "en", "es", "fr", "hi", "id", "it", "ja",
    "kk", "ko", "ky", "mr", "ms", "my", "pl", "pt", "ro", "ru", "sw", "te",
    "tg", "th", "tl", "tr", "uk", "ur", "uz", "vi", "zh",
}


def _language_subtags_compatible(a: str, b: str) -> bool:
    """True when two language codes plausibly name the same language once
    bridged across granularity — the shared logic behind both
    resolve_lang_code's and merge_lang_codes's "same language, differently
    spelled" bridging — while refusing a match that only "works" because
    an ISO-639 language code happens to spell the same two letters as an
    unrelated ISO-3166 country code (see INDEPENDENT_LANGUAGE_CODES above:
    Arabic "ar" must never match "es-ar" — Spanish, Argentina — just
    because Argentina's country code is also "AR").

    Always safe: the two codes share the same LANGUAGE subtag — "ko" and
    "ko-KR", or "es-ar" and "es-mx" by their common "es".

    Also safe, but only for a bare code that ISN'T itself a real,
    independent language (the agency's own country-code-style shorthand —
    "KZ" for Kazakh, "TJ" for Tajik, "BD" for Bengali, which aren't
    themselves recognized language codes) matching the REGION half of a
    fuller code ("KZ" against "kk-KZ"). A bare code that IS a real
    language in its own right is excluded from this side of the match
    entirely — it may only match by sharing an actual LANGUAGE subtag,
    never by coincidentally matching someone else's region."""
    a, b = a.strip().lower(), b.strip().lower()
    if _base_lang(a) == _base_lang(b):
        return True
    a_is_shorthand = "-" not in a and a not in INDEPENDENT_LANGUAGE_CODES
    b_is_shorthand = "-" not in b and b not in INDEPENDENT_LANGUAGE_CODES
    if a_is_shorthand and a in _subtags(b):
        return True
    if b_is_shorthand and b in _subtags(a):
        return True
    return False


def _freeze(value):
    """Makes a value hashable/comparable for the equality check in
    resolve_lang_code below — a document's row might be a plain string
    (e.g. Tone's register) or a dict of fields, depending on the doc."""
    if isinstance(value, dict):
        return tuple(sorted(value.items()))
    return value


def resolve_lang_code(requested: str, available, values: dict | None = None):
    """Matches a requested language code against a set/dict/iterable of
    codes actually present in one document, bridging granularity mismatches
    within that document — e.g. a target language selected as a plain "ko"
    still finds a region-qualified "ko-KR" row, and vice versa.

    Tries an exact (case-insensitive) match first. Failing that, falls
    back to matching the requested code against a compatible subtag of an
    available code (see _language_subtags_compatible — language part
    always, region part only for a genuine country-code-style shorthand,
    never for a bare code that's a real language in its own right) — but
    ONLY when exactly one available code is compatible; if several are
    (e.g. a document has both "es-ES" and "es-AR" with genuinely different
    rules for each), guessing would silently apply the wrong regional
    rule, so this returns None instead — the caller then treats the
    language as if it had no entry at all, exactly like today's "document
    uploaded but this language is missing" case, rather than picking one
    region at random.

    Some languages genuinely split into a handful of regional variants
    where every OTHER variant besides one or two special cases shares the
    same rule (e.g. Spanish: Spain and Argentina each have their own
    format, but "the rest of Latin America" is one shared format that can
    show up under any of several country codes — es-MX, es-CL, es-PE...).
    The agency already has a way to express that directly in the document
    itself: list every code that shares one row together in one cell,
    separated by "/" (already used for French: "fr-CI / fr-FR") — every
    code listed then resolves by exact match, no ambiguity at all. As a
    safety net for a code that WASN'T listed, if `values` (a {code: value}
    mapping) is passed, several same-base candidates still resolve when
    they all happen to carry the identical value — applying it is safe
    regardless of which one is picked, since they don't actually disagree.

    Returns the matching code from `available` (preserving its original
    casing), or None if nothing resolves safely.
    """
    requested = (requested or "").strip().lower()
    if not requested:
        return None

    avail_list = list(available)
    by_lower = {a.lower(): a for a in avail_list}
    if requested in by_lower:
        return by_lower[requested]

    matches = [a for a in avail_list if _language_subtags_compatible(requested, a)]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1 and values is not None:
        distinct = {_freeze(values[a]) for a in matches if a in values}
        if len(distinct) == 1:
            return matches[0]
    return None


def _lang_selected(lang: str, target_langs_filter: set[str]) -> bool:
    """Whether `lang` — a code straight from a FILE's own header, already
    normalized by _normalize_lang_label — should count as "selected" by a
    manager's target_langs_filter (the raw codes of whichever catalog
    checkboxes were ticked in the UI). A plain `in` check breaks the
    instant the file's own spelling doesn't EXACTLY match the catalog's
    spelling for the same language: a catalog still holding an older bare
    "pt" entry (from before Portuguese started defaulting to "pt-br") next
    to a freshly-uploaded file whose "PT" column now normalizes to
    "pt-br" is exactly the case that motivated this — Александр's
    detect-languages notice already told him "pt-br" would be checked
    (that endpoint has always bridged catalog/file spelling mismatches),
    so the actual check run must honor that instead of silently dropping
    the language over a spelling technicality.

    Deliberately narrower than resolve_lang_code itself: bridges ONLY when
    exactly one of the two sides is a bare code (no region at all) and the
    other is region-qualified — bare "pt" <-> file's "pt-br", or a
    country-code-style shorthand like "kz" <-> "kk-KZ" — via the same
    _language_subtags_compatible rules used everywhere else. Two codes
    that are BOTH already region-qualified are never bridged just because
    they happen to share a base language: "es-mx" and "es-es" are
    deliberately different, explicitly-added catalog languages (see
    merge_lang_codes), and ticking one must never silently sweep in the
    other's column too — resolve_lang_code's own base-language shortcut is
    too permissive for that case, so it isn't reused here. Refuses (rather
    than guessing) when a single FILE column could bridge to more than one
    ticked filter entry — e.g. a file's bare "es" column with both "es-ar"
    and "es-mx" ticked never silently ends up checked as just one of them.

    This is NOT symmetric with resolve_lang_code's own ambiguity refusal,
    and deliberately so: when the CATALOG side is the bare one instead
    (only a generic "es" ticked) and the file has several explicit
    columns for it ("es-ar" AND "es-mx" both present), each column bridges
    to that one bare entry independently and BOTH get selected — the
    manager's bare tick reads as "check Spanish, generically", and this
    module's guiding rule (see the PR/Peru and GEO catalog fixes elsewhere
    in this file) is that a language never silently drops out of a check
    just because of a granularity mismatch. Over-including a language the
    manager arguably meant to cover is a far smaller problem than the bug
    this function exists to fix (a language silently never checked at
    all), so this asymmetry is intentional, not a gap to close."""
    lang_low = lang.strip().lower()
    filt = {f.strip().lower() for f in target_langs_filter if f and f.strip()}
    if lang_low in filt:
        return True
    lang_is_bare = "-" not in lang_low
    matches = [
        f for f in filt
        if (("-" not in f) != lang_is_bare) and _language_subtags_compatible(lang_low, f)
    ]
    return len(matches) == 1


def merge_lang_codes(codes) -> list[str]:
    """Deduplicates a set of language codes gathered from several documents
    for display (e.g. the target-language picker) — when a base language
    has only one region variant across everything ("ko" in one doc, "ko-KR"
    in another), it's shown once, as its more specific spelling. When a
    base language genuinely has several distinct region variants (es-ES,
    es-AR, es-MX), each is kept as its own separate entry, since they mean
    different formatting rules and must be picked explicitly.

    Region-qualified codes are grouped by their first (language) subtag
    only — deliberately narrower than a full compatibility check, since
    two region-qualified codes should never merge just for sharing a
    region (hi-IN and mr-IN are different languages that happen to both be
    spoken in India). A BARE code with no region of its own is looser by
    nature — it's merged into whichever region-qualified group it's
    compatible with (see _language_subtags_compatible: a genuine
    country-code-style shorthand like Tone's "KZ" for Kazakh matches on
    ANY subtag, but a bare code that's a real independent language, like
    Arabic "ar", only matches by sharing an actual language subtag — it
    must never be absorbed into another language's group just because it
    happens to spell the same two letters as one of that group's REGIONS,
    e.g. Arabic "ar" vs Argentina's country code inside "es-ar"), and only
    when that's unambiguous (exactly one group matches)."""
    codes = [(c or "").strip() for c in codes if c and c.strip()]
    hyphenated = [c for c in codes if "-" in c]
    bare = [c for c in codes if "-" not in c]

    groups: dict[str, list[str]] = {}
    for code in hyphenated:
        groups.setdefault(_base_lang(code), []).append(code)

    for code in bare:
        low = code.lower()
        matching_bases = {
            base for base, variants in groups.items()
            if any(_language_subtags_compatible(low, v) for v in variants)
        }
        if len(matching_bases) == 1:
            groups[next(iter(matching_bases))].append(code)
        else:
            groups.setdefault(low, []).append(code)

    result = []
    for base, variants in groups.items():
        distinct = sorted(set(variants))
        regioned = [v for v in distinct if "-" in v]
        if len(regioned) >= 2:
            # Several genuinely different regional variants (es-ES vs
            # es-AR vs es-MX) — keep each; drop any bare/generic spelling
            # of the same base, since it's ambiguous which region it means.
            result.extend(regioned)
        elif len(regioned) == 1:
            # Only one spelling actually matters for this language — a
            # bare "ko" alongside it is the same language, not a second one.
            result.append(regioned[0])
        else:
            result.append(distinct[0])
    return sorted(result)


def _find_header_row(ws, max_scan: int = 5) -> int:
    best_row, best_score = 1, -1
    for r in range(1, min(max_scan, ws.max_row) + 1):
        score = 0
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and LANG_CODE_RE.match(v.strip().lower()):
                score += 1
        if score > best_score:
            best_row, best_score = r, score
    return best_row


def parse_workbook(file_bytes: bytes) -> list[dict]:
    """Returns a list of parsed sheets: each with lang codes found and rows."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheets = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue

        header_row = _find_header_row(ws)
        context_col = max_length_col = None
        lang_cols: dict[int, str] = {}
        unrecognized: list[str] = []

        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if not isinstance(v, str) or not v.strip():
                continue
            label = v.strip()
            if _is_context_col(label):
                context_col = c
            elif _is_max_length_col(label):
                max_length_col = c
            elif _is_meta_col(label):
                continue
            elif _is_limit_spec_col(label):
                # Obviously a limit/spec column, not a near-miss language
                # code — drop it silently instead of flagging it to the
                # manager as "not recognized as a language".
                continue
            else:
                normalized = _normalize_lang_label(label)
                if " " in normalized or len(normalized) > 12:
                    # Looks like prose, not a language code — skip rather
                    # than misread a stray comment column as a "language".
                    unrecognized.append(label)
                else:
                    lang_cols[c] = normalized

        if not lang_cols:
            continue

        rows = []
        last_context = ""
        for r in range(header_row + 1, ws.max_row + 1):
            values = {code: ws.cell(row=r, column=c).value for c, code in lang_cols.items()}
            if all(v is None or str(v).strip() == "" for v in values.values()):
                continue

            context = ""
            if context_col:
                raw_ctx = ws.cell(row=r, column=context_col).value
                context = str(raw_ctx).strip() if raw_ctx else ""
            if context:
                last_context = context
            else:
                context = last_context

            max_length = None
            if max_length_col:
                raw_ml = ws.cell(row=r, column=max_length_col).value
                if isinstance(raw_ml, (int, float)):
                    max_length = int(raw_ml)

            rows.append({
                "excel_row": r,
                "context": context,
                "max_length": max_length,
                "values": {code: ("" if v is None else str(v)) for code, v in values.items()},
            })

        sheets.append({
            "sheet_name": sheet_name,
            "languages": sorted(lang_cols.values()),
            "rows": rows,
            "unrecognized_columns": unrecognized,
        })

    return sheets


def pick_source_lang(sheets: list[dict], preferred: str | None) -> str:
    """Resolves the manager's chosen source language (e.g. "ru", from the
    RU/EN buttons in the UI) against the language codes actually found as
    column headers in the uploaded file. Uses resolve_lang_code rather than
    a literal match, since the file's own column can be a differently
    granular spelling of the same language (e.g. "ru-RU" for a plain "ru")
    — Александр hit this: his file's Russian column wasn't spelled exactly
    "ru", the literal check silently missed it, and the source language
    silently fell back to English (whatever column happened to be labeled
    "en-001"), even though he'd picked Russian. Falls back to English, then
    alphabetically first, only when the requested language truly isn't in
    the file at all."""
    all_langs: set[str] = set()
    for s in sheets:
        all_langs.update(s["languages"])
    if preferred:
        resolved = resolve_lang_code(preferred, all_langs)
        if resolved:
            return resolved
    if "en" in all_langs:
        return "en"
    return sorted(all_langs)[0] if all_langs else "en"


async def _check_language_for_sheet(
    sheet: dict,
    lang: str,
    source_lang: str,
    checks: list[str],
    extra_instructions: str,
    semaphore: asyncio.Semaphore,
    tone_register: str = "",
) -> tuple[list[dict], float]:
    relevant_rows = []
    ai_items = []
    for row in sheet["rows"]:
        src = row["values"].get(source_lang, "")
        tgt = row["values"].get(lang, "")
        if not src.strip() and not tgt.strip():
            continue
        if _is_do_not_translate(tgt):
            continue
        relevant_rows.append(row)
        ai_items.append({"context": row["context"], "source": src, "translation": tgt})

    if not relevant_rows:
        return [], 0.0

    async with semaphore:
        ai_findings_by_idx, cost_usd, truncated = await run_ai_checks_batch(
            ai_items, checks, extra_instructions, tone_register, lang, source_lang
        )

    out = []
    for idx, row in enumerate(relevant_rows):
        src = row["values"].get(source_lang, "")
        tgt = row["values"].get(lang, "")
        findings = run_rule_checks(src, tgt, checks, max_length=row["max_length"], lang_code=lang)
        findings += ai_findings_by_idx.get(idx, [])
        if findings:
            out.append({
                "excel_row": row["excel_row"],
                "context": row["context"],
                "source": src,
                "translation": tgt,
                "findings": findings,
            })
    if truncated:
        # The model's response for this language got cut off mid-array —
        # some rows may never have been checked by it at all. Surfaced as
        # its own synthetic entry rather than silently showing whatever
        # partial findings survived as if they were the complete picture
        # (see Александр's "incomplete report" on a large Spanish upload).
        out.append({
            "excel_row": 0,
            "context": "⚠ Системное предупреждение",
            "source": "",
            "translation": "",
            "findings": [_truncation_warning()],
        })
    return out, cost_usd


async def run_multi_check(
    sheets: list[dict],
    source_lang: str,
    checks: list[str],
    extra_instructions: str = "",
    tone_for_lang=None,
    target_langs_filter: set[str] | None = None,
) -> dict:
    """
    tone_for_lang: a callable(lang_code) -> "formal"/"informal"/"" (or ""
    if nothing for that language), already narrowed to just what this one
    target language needs — see app.project_docs. Each target language gets
    its own call, so the AI prompt for e.g. "es-mx" never carries the other
    34 languages' rows.

    target_langs_filter: when given, only these languages are checked even
    if the file has more columns — lets a manager check a subset of a
    large upload instead of every language every time.
    """
    tone_for_lang = tone_for_lang or (lambda lang: "")
    semaphore = asyncio.Semaphore(AI_CONCURRENCY)
    result_sheets = []
    total_findings = 0
    total_rows_checked = 0
    total_cost_usd = 0.0

    for sheet in sheets:
        target_langs = [l for l in sheet["languages"] if l != source_lang]
        if target_langs_filter is not None:
            target_langs = [l for l in target_langs if _lang_selected(l, target_langs_filter)]
        tasks = [
            _check_language_for_sheet(
                sheet, lang, source_lang, checks, extra_instructions, semaphore,
                tone_for_lang(lang),
            )
            for lang in target_langs
        ]
        per_lang_results = await asyncio.gather(*tasks) if tasks else []

        languages_out = {}
        for lang, (findings_list, lang_cost) in zip(target_langs, per_lang_results):
            languages_out[lang] = findings_list
            total_findings += sum(len(f["findings"]) for f in findings_list)
            total_cost_usd += lang_cost

        total_rows_checked += len(sheet["rows"])
        result_sheets.append({
            "sheet_name": sheet["sheet_name"],
            "source_lang": source_lang,
            "languages_checked": target_langs,
            "languages": languages_out,
            "unrecognized_columns": sheet.get("unrecognized_columns", []),
        })

    summary = {
        "sheets": len(result_sheets),
        "rows_checked": total_rows_checked,
        "languages_checked": sorted({l for s in result_sheets for l in s["languages_checked"]}),
        "total_findings": total_findings,
        "cost_usd": total_cost_usd,
    }
    return {"sheets": result_sheets, "summary": summary}


# ------------------------------------------------- large jobs: batch mode ---
# See BATCH_THRESHOLD_CHARS above. Instead of awaiting every language's AI
# call directly (run_multi_check), a large job is prepared as a "skeleton"
# (rule-based findings, computed instantly and for free) plus one Anthropic
# Message Batch request per language; once that batch finishes — polled from
# app.main — finalize_batch_results merges the AI findings back in to
# produce the exact same {"sheets": [...], "summary": {...}} shape as
# run_multi_check, so the frontend doesn't need to know which path ran.

def estimate_check_volume(
    sheets: list[dict], source_lang: str, target_langs_filter: set[str] | None = None
) -> int:
    """Rough proxy (total characters, summed across every checkable row ×
    every target language) for how expensive/slow a synchronous run would
    be. Not exact — a ballpark is all that's needed to pick a processing
    mode. Respects target_langs_filter so checking only a handful of a
    file's languages doesn't get pushed into the slow queue on the
    strength of languages that won't even be checked this run."""
    total = 0
    for sheet in sheets:
        target_langs = [l for l in sheet["languages"] if l != source_lang]
        if target_langs_filter is not None:
            target_langs = [l for l in target_langs if _lang_selected(l, target_langs_filter)]
        for row in sheet["rows"]:
            src = row["values"].get(source_lang, "")
            for lang in target_langs:
                tgt = row["values"].get(lang, "")
                if not src.strip() and not tgt.strip():
                    continue
                if _is_do_not_translate(tgt):
                    continue
                total += len(src) + len(tgt)
    return total


def build_batch_plan(
    sheets: list[dict],
    source_lang: str,
    checks: list[str],
    extra_instructions: str = "",
    tone_for_lang=None,
    target_langs_filter: set[str] | None = None,
) -> tuple[list[dict], dict]:
    """Prepares everything needed to submit one Anthropic Message Batch
    covering every (sheet, target language) pair in this upload, plus a
    JSON-serializable "skeleton" — already-computed rule-based findings —
    to merge the AI results into later via finalize_batch_results.

    Returns (batch_requests, skeleton). batch_requests is a list of
    {"custom_id": str, "prompt": str} ready for
    claude_client.create_message_batch; it can be empty if no AI check
    types were selected at all, in which case there's nothing to submit and
    finalize_batch_results(skeleton, {}) is already the final answer.
    """
    tone_for_lang = tone_for_lang or (lambda lang: "")
    requests: list[dict] = []
    skeleton_sheets = []

    for s_idx, sheet in enumerate(sheets):
        target_langs = [l for l in sheet["languages"] if l != source_lang]
        if target_langs_filter is not None:
            target_langs = [l for l in target_langs if _lang_selected(l, target_langs_filter)]
        languages_skeleton = {}

        for lang_idx, lang in enumerate(target_langs):
            relevant_rows = []
            ai_items = []
            for row in sheet["rows"]:
                src = row["values"].get(source_lang, "")
                tgt = row["values"].get(lang, "")
                if not src.strip() and not tgt.strip():
                    continue
                if _is_do_not_translate(tgt):
                    continue
                relevant_rows.append(row)
                ai_items.append({"context": row["context"], "source": src, "translation": tgt})

            # Rule-based findings are free and instant — compute them now
            # rather than waiting on the batch for them too.
            base_rows = []
            for row in relevant_rows:
                src = row["values"].get(source_lang, "")
                tgt = row["values"].get(lang, "")
                findings = run_rule_checks(src, tgt, checks, max_length=row["max_length"], lang_code=lang)
                base_rows.append({
                    "excel_row": row["excel_row"],
                    "context": row["context"],
                    "source": src,
                    "translation": tgt,
                    "findings": findings,
                })

            # Built from the sheet/position index only, never from `lang`
            # itself — Anthropic's Batches API requires custom_id to match
            # ^[a-zA-Z0-9_-]{1,64}$, but a language code comes straight from
            # a column header in whatever file gets uploaded and can't be
            # trusted to satisfy that (e.g. a Cyrillic character that looks
            # identical to a Latin one, from a copy-pasted "fr-CI" header,
            # is enough to make Anthropic reject the WHOLE batch — every
            # language in it, not just the bad one — with a 400). Rebuilt
            # this way, custom_id is always safe regardless of what's in
            # the file.
            custom_id = f"s{s_idx}-t{lang_idx}"
            model = _model_for_lang(lang)
            prompt, number_to_index = build_batch_prompt(
                ai_items, checks, extra_instructions,
                tone_for_lang(lang), lang, source_lang,
            )
            if prompt is not None:
                requests.append({"custom_id": custom_id, "prompt": prompt, "model": model})

            languages_skeleton[lang] = {
                "custom_id": custom_id if prompt is not None else None,
                # Needed later by finalize_batch_results to price this
                # language's usage at the right per-token rate.
                "model": model,
                "number_to_index": {str(k): v for k, v in number_to_index.items()},
                "rows": base_rows,
            }

        skeleton_sheets.append({
            "sheet_name": sheet["sheet_name"],
            "target_langs": target_langs,
            "languages": languages_skeleton,
            "unrecognized_columns": sheet.get("unrecognized_columns", []),
            "row_count": len(sheet["rows"]),
        })

    # checks rides along in the skeleton so finalize_batch_results (called
    # later, sometimes in a completely different request once the
    # Anthropic batch has ended) can still filter the model's response to
    # only what was actually asked for.
    skeleton = {"sheets": skeleton_sheets, "source_lang": source_lang, "checks": checks}
    return requests, skeleton


def finalize_batch_results(skeleton: dict, ai_results_by_custom_id: dict[str, dict]) -> dict:
    """Merges AI findings (once the Anthropic batch has ended) into the
    rule-based skeleton from build_batch_plan, producing the same
    {"sheets": [...], "summary": {...}} shape run_multi_check returns.

    ai_results_by_custom_id: {custom_id: {"text": str | None, "usage": dict}}
    — see claude_client.get_batch_results. Missing/empty entries (e.g. no
    batch was actually submitted) simply contribute no findings and no cost."""
    result_sheets = []
    total_findings = 0
    total_rows_checked = 0
    total_cost_usd = 0.0

    for sheet in skeleton["sheets"]:
        languages_out = {}
        for lang, lang_skel in sheet["languages"].items():
            ai_grouped: dict[int, list[dict]] = {}
            warning_finding = None
            custom_id = lang_skel["custom_id"]
            if custom_id is not None:
                ai_result = ai_results_by_custom_id.get(custom_id)
                if ai_result is None:
                    # Expected result never showed up in the batch's .jsonl
                    # at all — same class of silent data loss as an
                    # errored/expired request, so it gets the same warning
                    # rather than quietly counting as "nothing found".
                    warning_finding = _ai_failure_warning("результат не получен")
                elif ai_result.get("result_type") != "succeeded":
                    warning_finding = _ai_failure_warning(ai_result.get("result_type") or "неизвестная ошибка")
                else:
                    raw = parse_json_array(ai_result.get("text"))
                    # JSON round-trips dict keys as strings — restore int keys.
                    number_to_index = {int(k): v for k, v in lang_skel["number_to_index"].items()}
                    ai_grouped = group_batch_findings(raw, number_to_index)
                    # Guarantees the model's response never smuggles in a check
                    # type the manager didn't ask for, even if it ignored the
                    # prompt's instruction to stick to the requested list.
                    ai_grouped = {
                        idx: _filter_findings_by_checks(fs, skeleton.get("checks", []))
                        for idx, fs in ai_grouped.items()
                    }
                    if ai_result.get("stop_reason") == "max_tokens":
                        # Response got cut off mid-array — some rows may
                        # never have been checked by the AI at all (see
                        # Александр's "incomplete Spanish report").
                        warning_finding = _truncation_warning()
                if ai_result is not None:
                    total_cost_usd += _usage_cost(lang_skel.get("model", ""), ai_result.get("usage"), batch=True)

            findings_list = []
            for idx, row in enumerate(lang_skel["rows"]):
                findings = list(row["findings"]) + ai_grouped.get(idx, [])
                if findings:
                    findings_list.append({
                        "excel_row": row["excel_row"],
                        "context": row["context"],
                        "source": row["source"],
                        "translation": row["translation"],
                        "findings": findings,
                    })
            if warning_finding is not None:
                findings_list.append({
                    "excel_row": 0,
                    "context": "⚠ Системное предупреждение",
                    "source": "",
                    "translation": "",
                    "findings": [warning_finding],
                })
            languages_out[lang] = findings_list
            total_findings += sum(len(f["findings"]) for f in findings_list)

        total_rows_checked += sheet["row_count"]
        result_sheets.append({
            "sheet_name": sheet["sheet_name"],
            "source_lang": skeleton["source_lang"],
            "languages_checked": sheet["target_langs"],
            "languages": languages_out,
            "unrecognized_columns": sheet["unrecognized_columns"],
        })

    summary = {
        "sheets": len(result_sheets),
        "rows_checked": total_rows_checked,
        "languages_checked": sorted({l for s in result_sheets for l in s["languages_checked"]}),
        "total_findings": total_findings,
        "cost_usd": total_cost_usd,
    }
    return {"sheets": result_sheets, "summary": summary}


async def submit_multi_check_batch(requests: list[dict]) -> str | None:
    return await create_message_batch(requests)


async def cancel_multi_check_batch(batch_id: str) -> None:
    """Best-effort cancel for a manager who no longer wants to wait for (or
    pay for) a still-processing upload — e.g. they want to switch to
    "Срочно" instead, or simply changed their mind. Anthropic may reject
    this (most likely because the batch had already ended right as the
    manager clicked cancel) — that's fine, the caller is deleting its own
    record either way, so a failure here should never block that."""
    try:
        await cancel_message_batch(batch_id)
    except Exception:
        pass


def _batch_progress(status: dict) -> dict:
    """Turns Anthropic's own request_counts ({"processing": n, "succeeded":
    n, "errored": n, "canceled": n, "expired": n}) into a simple {"done",
    "total"} the UI can show as a real progress readout. Anthropic doesn't
    publish an ETA for a batch job, so a made-up time estimate would just be
    a guess — this is the one number we actually know is true."""
    counts = status.get("request_counts") or {}
    total = sum(counts.values())
    done = total - counts.get("processing", 0)
    return {"done": done, "total": total}


async def try_finalize_batch(batch_id: str, skeleton: dict) -> tuple[dict | None, dict]:
    """Returns (finalized_results, progress). finalized_results is the
    completed results dict once the Anthropic batch has ended, otherwise
    None (still processing — caller should try again later). progress is
    always {"done": int, "total": int} from Anthropic's request_counts, so
    the caller can surface real progress even while still waiting."""
    status = await get_batch_status(batch_id)
    progress = _batch_progress(status)
    if status.get("processing_status") != "ended":
        return None, progress
    results_url = status.get("results_url")
    ai_results_by_custom_id = await get_batch_results(results_url) if results_url else {}
    return finalize_batch_results(skeleton, ai_results_by_custom_id), progress


def _plural_ru(n: int, one: str, few: str, many: str) -> str:
    """Standard Russian count-noun pluralization — mirrors the frontend's
    pluralRu (lang.ts), kept as a separate copy since this side is Python."""
    mod10, mod100 = n % 10, n % 100
    if 11 <= mod100 <= 14:
        return many
    if mod10 == 1:
        return one
    if 2 <= mod10 <= 4:
        return few
    return many


def _format_minutes_ru(minutes: float) -> str:
    """Mirrors the frontend's formatElapsedMinutesRu (lang.ts) — used in the
    downloadable report's summary line, below."""
    whole = round(minutes)
    if whole < 1:
        return "меньше минуты"
    return f"{whole} {_plural_ru(whole, 'минута', 'минуты', 'минут')}"


def build_report_workbook(
    filename: str, source_lang: str, results: dict, duration_minutes: float | None = None
) -> bytes:
    """Builds a downloadable .xlsx with one row per finding. duration_minutes
    (how long the check itself took, end to end) is optional — omitted
    entirely for an older record that predates this being tracked, rather
    than showing a misleading "0 минут"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QA Findings"
    if duration_minutes is not None:
        ws.append([f"Проверка «{filename}» ({source_lang}) — заняла {_format_minutes_ru(duration_minutes)}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
        ws.cell(row=1, column=1).font = openpyxl.styles.Font(bold=True)
        ws.append([])  # spacer row before the header
    ws.append(["Лист", "Строка в файле", "Контекст", "Язык", "Серьёзность", "Тип", "Проблема", "Источник", "Перевод"])
    # Captured AFTER the append above (not computed from the pre-append
    # max_row) — an appended blank spacer row still advances openpyxl's
    # internal row cursor even though it holds no cells, so computing this
    # beforehand pointed one row too early and left the header itself out
    # of the filter/freeze range below.
    header_row = ws.max_row
    for col_idx, width in enumerate([18, 14, 28, 8, 12, 14, 50, 40, 40], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    for sheet in results.get("sheets", []):
        for lang, findings_list in sheet.get("languages", {}).items():
            for item in findings_list:
                for f in item["findings"]:
                    ws.append([
                        sheet["sheet_name"],
                        item["excel_row"],
                        item["context"],
                        lang,
                        f.get("severity", ""),
                        f.get("type", ""),
                        f.get("message", ""),
                        item["source"],
                        item["translation"],
                    ])

    # Turns on Excel's own column filter dropdowns on the header row — lets
    # Александр filter to just one "Язык" (or any other column) using the
    # filter control Excel already gives him, defaulting to showing
    # everything exactly as it does today. Only makes sense once there's at
    # least one data row below the header.
    if ws.max_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:I{ws.max_row}"
        ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
